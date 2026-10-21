#!/usr/bin/env python3
"""
Disfluency Annotator - Python port of annotate_disfluencies_video.m (v4.1)

GUI for annotating stuttering-like disfluencies (SLD), typical disfluencies
(TD) and speech errors from video.

    python disfluency_annotator.py                 # open, then "Load video"
    python disfluency_annotator.py path/to/video   # open and load a video

Video is decoded with FFmpeg (PyAV): frame-accurate, full resolution, any
codec FFmpeg reads (iPhone .MOV / HEVC included), rotation metadata honoured,
scaled to your screen's real pixels with high-quality (Lanczos / smooth)
filtering. No "smooth copy" conversion is needed.

Event times come from the AUDIO sample clock, exactly as in the MATLAB version,
and annotation files are read / written in the same format (events, codes,
annotator_data sheets), so old files load here and new files load there.

Workflow (example "the t-t-table is on the ground"):
  1. Highlight speech, DOUBLE-CLICK it, type the sentence, Enter.
  2. Highlight "table" (or click its piece) + Enter = fluent event selected.
  3. Highlight "t-t" (or one "t-") + Enter, tick types (R P B T / 1-6), Enter.
     W = whole episode. Esc = done with this word.
Space plays the highlight (L loop, 1-4 speed). Ctrl/Cmd+S saves. F1 = help.
"""

import sys
import os
import re
import math
import time
import csv
import bisect
import threading
from collections import deque
from dataclasses import dataclass, replace
from types import SimpleNamespace

import numpy as np

try:
    import av
except ImportError:
    sys.exit('PyAV is missing:  pip install av')

try:
    import sounddevice as sd
except Exception:            # PortAudio not found etc. -> video plays silently
    sd = None

import pyqtgraph as pg
from PySide6 import QtCore, QtGui, QtWidgets
from PySide6.QtCore import Qt, QPointF, QRectF, QTimer, QSettings, QEvent
from PySide6.QtGui import QImage, QPainter, QColor, QPen, QBrush, QFont, QFontMetrics

pg.setConfigOptions(imageAxisOrder='row-major', antialias=False)

# ----------------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------------
APPNAME = 'Disfluency Annotator'
APPVER = '5.0 (Python)'

DISFL = [
    dict(abbr='rep',  key='r', name='repetition (SLD)',   color=(0.85, 0.15, 0.15)),
    dict(abbr='prol', key='p', name='prolongation (SLD)', color=(0.10, 0.60, 0.20)),
    dict(abbr='bl',   key='b', name='block (SLD)',        color=(0.15, 0.35, 0.90)),
    dict(abbr='td',   key='t', name='typical disfluency', color=(0.95, 0.78, 0.05)),
]
ERRS = [
    dict(abbr='er-os', name='omitted sound'),
    dict(abbr='er-ow', name='omitted word'),
    dict(abbr='er-is', name='inserted sound'),
    dict(abbr='er-iw', name='word insertion'),
    dict(abbr='er-t',  name='transposition'),
    dict(abbr='er-s',  name='substitution'),
]
ERR_KEY = 'e'
FLU_KEY = 'f'
FLUENT_COLOR = (0.78, 0.78, 0.78)
ERR_COLOR = (0.85, 0.00, 0.85)
TRANS_COLOR = (0.45, 0.45, 0.45)

# False (per spec): disfluency / error rows are saved WITHOUT a transcript.
DISFL_HAS_TRANSCRIPT = False

IPA_ROWS = ['θðʃʒxŋɹʔ',
            'ɛæɪəʊʌɑɔɚɝ']

BG = (0.94, 0.94, 0.94)
SEL_COLOR = (0.45, 0.45, 0.45)
OV_COLOR = (0.20, 0.45, 0.95)
CUR_COLOR = (0.90, 0.50, 0.00)
PLAY_COLOR = (0.00, 0.55, 0.55)
SPEEDS = [1, 0.75, 0.5, 0.25]
SPEED_NAMES = ['1x', '0.75x', '0.5x', '0.25x']

MINVIEW = 0.02
MINSEL = 0.005
EDGE_PX = 8
MAXUNDO = 40
TOL = 1e-6
CACHE_MAX_S = 60       # longest highlight kept in memory for smooth looping (s)
CACHE_MAX_MB = 1200    # memory cap for that (MB)
NAN = float('nan')


def qc(rgb, a=1.0):
    r, g, b = [min(1.0, max(0.0, float(v))) for v in rgb]
    return QColor.fromRgbF(r, g, b, a)


def mix(rgb, w, base):
    return tuple(base + w * c for c in rgb)


def isnan(x):
    return x is None or (isinstance(x, float) and math.isnan(x))


# ----------------------------------------------------------------------------
# Data model
# ----------------------------------------------------------------------------
@dataclass
class Ev:
    start: float
    end: float
    kind: str              # 'trans', 'fluent', 'disfl', 'error'
    cat: str = ''
    transcript: str = ''
    notes: str = ''
    uid: int = 0
    link: int = 0          # uid of the fluent word a disfl / error resolves to
    wi: int = 0            # word index (older files)
    tok: int = 0           # uid of the transcript piece the event sits on
    gStart: float = NAN
    gStop: float = NAN


def fmt_time(t):
    if t is None or isnan(t):
        return '--:--.---'
    m = int(math.floor(t / 60))
    return '%02d:%06.3f' % (m, t - 60 * m)


def parse_time(s):
    s = str(s).strip().lower().replace(' ', '')
    if not s:
        return NAN
    s = re.sub(r'(sec|secs|seconds|s)$', '', s)
    if 'm' in s:
        s = re.sub(r'(min|mins|minutes|m)', ':', s)
        if s.endswith(':'):
            s += '0'
    try:
        v = [float(p) for p in s.split(':')]
    except ValueError:
        return NAN
    if len(v) > 3 or any(x < 0 for x in v):
        return NAN
    t = 0.0
    for x in v:
        t = t * 60 + x
    return t


def token_word(s):
    s = str(s).strip()
    w = re.sub(r'^[-.,;:!?"\'()]+|[-.,;:!?"\'()]+$', '', s)
    return w if w else s


def split_tokens(s):
    toks = re.findall(r'[^\s-]+-*', str(s))
    return toks if toks else [str(s).strip()]


def norm_word(s):
    w = str(s).strip().lower()
    return re.sub(r'^[.,;:!?"\'()\[\]-]+|[.,;:!?"\'()\[\]-]+$', '', w)


def norm_word_keep(tok):
    w = re.sub(r'^[.,;:!?"\'()\[\]-]+|[.,;:!?"\'()\[\]-]+$', '', str(tok).strip())
    return w if w else str(tok)


def to_num(v):
    if v is None:
        return NAN
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return NAN


def to_str(v):
    if v is None:
        return ''
    if isinstance(v, float):
        if math.isnan(v):
            return ''
        return ('%d' % v) if v.is_integer() else repr(v)
    s = str(v).strip()
    return '' if s.lower() in ('nan', '<missing>', 'none') else s


def one_line(s):
    return str(s).replace('\n', ' ').replace('\r', ' ')


# ----------------------------------------------------------------------------
# Media: audio, video decoding, frame preparation
# ----------------------------------------------------------------------------
def load_audio(path):
    """Decode the audio track to mono float32. Returns (audio, fs, origin_s)."""
    c = av.open(path)
    try:
        if not c.streams.audio:
            raise RuntimeError('This video has no audio track, so it cannot be annotated here.')
        a = c.streams.audio[0]
        fs = int(a.rate or a.codec_context.sample_rate)
        origin = float(a.start_time * a.time_base) if a.start_time is not None else 0.0
        rs = av.AudioResampler(format='flt', layout='mono', rate=fs)
        chunks = []
        for fr in c.decode(a):
            for o in rs.resample(fr):
                chunks.append(o.to_ndarray().reshape(-1))
        try:
            for o in rs.resample(None):
                chunks.append(o.to_ndarray().reshape(-1))
        except Exception:
            pass
        if not chunks:
            raise RuntimeError('The audio track is empty.')
        audio = np.ascontiguousarray(np.concatenate(chunks).astype(np.float32))
        peak = float(np.max(np.abs(audio))) if len(audio) else 0.0
        if peak > 0.98:                       # loud recording: scale down instead of clipping
            audio *= np.float32(0.98 / peak)
        return audio, fs, origin
    finally:
        c.close()


class VideoSource:
    """Frame-accurate access to a video stream (one per thread)."""

    def __init__(self, path, origin):
        self.c = av.open(path)
        self.s = self.c.streams.video[0]
        try:
            self.s.thread_type = 'AUTO'
        except Exception:
            pass
        self.tb = float(self.s.time_base)
        r = self.s.average_rate or self.s.guessed_rate or self.s.base_rate
        self.fps = float(r) if r else 30.0
        if not (1 <= self.fps <= 480):
            self.fps = 30.0
        self.origin = origin
        self.W = self.s.codec_context.width
        self.H = self.s.codec_context.height
        self._it = None
        self.cur = None
        self.pending = None
        self.seek(0.0)

    def close(self):
        try:
            self.c.close()
        except Exception:
            pass

    def seek(self, t):
        ts = int(math.floor((t + self.origin) / self.tb)) - 1
        try:
            self.c.seek(max(ts, 0), stream=self.s, backward=True, any_frame=False)
        except Exception:
            try:
                self.c.seek(0, stream=self.s)
            except Exception:
                pass
        self._it = self.c.decode(self.s)
        self.cur = None
        self.pending = None

    def _ftime(self, f):
        if f.pts is not None:
            return f.pts * self.tb - self.origin
        if f.time is not None:
            return f.time - self.origin
        return None

    def pull(self):
        """Next decoded frame as (frame, t) or None at the end."""
        if self.pending is not None:
            x = self.pending
            self.pending = None
            return x
        bad = 0
        while True:
            try:
                f = next(self._it)
            except StopIteration:
                return None
            except av.error.EOFError:
                return None
            except Exception:
                bad += 1
                if bad > 50:
                    return None
                continue
            t = self._ftime(f)
            if t is None:
                continue
            return (f, t)

    def frame_at(self, t):
        """The frame on screen at time t: (frame, frame_time)."""
        if self.cur is None or t < self.cur[1] - 1e-4 or t > self.cur[1] + 1.0:
            self.seek(t)
        while True:
            x = self.pull()
            if x is None:
                break
            if x[1] > t + 1e-4:
                self.pending = x
                if self.cur is None:
                    self.cur = x
                break
            self.cur = x
        return self.cur


def frame_rgb(f, w=None, h=None):
    """av.VideoFrame -> contiguous RGB array, optionally scaled with Lanczos."""
    if w:
        try:
            arr = f.to_ndarray(width=w, height=h, format='rgb24', interpolation='LANCZOS')
        except Exception:
            try:
                arr = f.reformat(width=w, height=h, format='rgb24').to_ndarray()
            except Exception:
                arr = f.to_ndarray(format='rgb24')
    else:
        arr = f.to_ndarray(format='rgb24')
    return np.ascontiguousarray(arr)


def frame_rotation(f):
    try:
        r = getattr(f, 'rotation', 0) or 0
        return int(round(float(r) / 90.0)) % 4
    except Exception:
        return 0


class Prepared:
    __slots__ = ('t', 'img', 'arr', 'cover', 'gen')


def _qimage(arr):
    h, w = arr.shape[:2]
    return QImage(arr.data, w, h, arr.strides[0], QImage.Format.Format_RGB888)


def prepare(src, t, P):
    """Make a display-ready image of frame src (av.VideoFrame or raw RGB array)
    for the current zoom and the panel's REAL pixel size (hi-dpi aware)."""
    is_frame = not isinstance(src, np.ndarray)
    k, zoom, tw, th, W, H = P.rot, P.zoom, max(P.tw, 16), max(P.th, 16), P.W, P.H
    p = Prepared()
    p.t = t
    p.gen = P.gen
    if zoom is None:
        s = min(tw / W, th / H, 1.0)
        ow, oh = max(2, int(round(W * s))), max(2, int(round(H * s)))
        if is_frame:
            rw, rh = (oh, ow) if k % 2 else (ow, oh)
            arr = frame_rgb(src, rw, rh) if s < 1 else frame_rgb(src)
            if k:
                arr = np.ascontiguousarray(np.rot90(arr, k))
            img = _qimage(arr)
        else:
            arr = np.ascontiguousarray(np.rot90(src, k)) if k else src
            img = _qimage(arr)
            if s < 1:
                img = img.scaled(ow, oh, Qt.IgnoreAspectRatio, Qt.SmoothTransformation)
                arr = None
        p.cover = (0, 0, W, H)
    else:
        raw = frame_rgb(src) if is_frame else src
        full = np.rot90(raw, k) if k else raw
        x0, y0, w, h = zoom
        arr = np.ascontiguousarray(full[y0:y0 + h, x0:x0 + w])
        img = _qimage(arr)
        s = min(tw / w, th / h)
        if s < 1:
            img = img.scaled(max(2, int(w * s)), max(2, int(h * s)),
                             Qt.IgnoreAspectRatio, Qt.SmoothTransformation)
            arr = None
        p.cover = tuple(zoom)
    if arr is None:
        img = img.copy()
    p.img = img
    p.arr = arr
    return p


class ViewParams:
    """Zoom / rotation / display size shared with the decoding thread."""

    def __init__(self, W, H, rot):
        self.lock = threading.Lock()
        self.W0, self.H0 = W, H
        self.rot = rot
        self.zoom = None
        self.tw, self.th = 1280, 720
        self.gen = 0

    def dims(self):
        return (self.H0, self.W0) if self.rot % 2 else (self.W0, self.H0)

    def update(self, **kw):
        with self.lock:
            for k, v in kw.items():
                setattr(self, k, v)
            self.gen += 1

    def snapshot(self):
        with self.lock:
            W, H = self.dims()
            return SimpleNamespace(rot=self.rot, zoom=self.zoom, tw=self.tw, th=self.th,
                                   W=W, H=H, gen=self.gen)


class PlaybackWorker(threading.Thread):
    """Decodes ahead of the audio clock in the background; caches a played
    highlight in memory so loops replay without decoding."""

    def __init__(self, path, origin, params, fps):
        super().__init__(daemon=True)
        self.path, self.origin, self.params, self.fps = path, origin, params, fps
        self.cv = threading.Condition()
        self.active = False
        self.quit = False
        self.target = 0.0
        self.req = None
        self.buf = deque()
        self.cache = None
        self.cache_range = None

    # -- called from the GUI --
    def begin(self, t0, cache_range):
        with self.cv:
            self.active = True
            self.target = t0
            self.cache_range = cache_range
            self.req = (t0, cache_range)
            self.cv.notify()

    def pause(self):
        with self.cv:
            self.active = False

    def set_target(self, t):
        self.target = t

    def invalidate(self):
        with self.cv:
            self.buf.clear()
            self.cache = None
            if self.active:
                self.req = (self.target, self.cache_range)
            self.cv.notify()

    def stop(self):
        with self.cv:
            self.quit = True
            self.cv.notify()

    def _cache_hit(self, t):
        c = self.cache
        if c is None or not c['complete'] or not c['times'] or c['gen'] != self.params.gen:
            return None
        if c['times'][0] - 1e-3 <= t <= c['times'][-1] + 1.0 / self.fps:
            i = bisect.bisect_right(c['times'], t + 1e-4) - 1
            return c['frames'][max(i, 0)]
        return None

    def get(self, t):
        with self.cv:
            p = self._cache_hit(t)
            if p is not None:
                return p
            best = None
            for q in self.buf:
                if q.t <= t + 1e-4:
                    best = q
                else:
                    break
            return best

    def cached(self, t):
        with self.cv:
            return self._cache_hit(t)

    # -- thread --
    def run(self):
        try:
            src = VideoSource(self.path, self.origin)
        except Exception:
            return
        fp = 1.0 / self.fps
        seek_t = 0.0
        while True:
            with self.cv:
                while not self.quit and not self.active and self.req is None:
                    self.cv.wait(0.2)
                if self.quit:
                    break
                req, self.req = self.req, None
                tgt = self.target
            if req is not None:
                t0, cr = req
                P = self.params.snapshot()
                with self.cv:
                    self.buf.clear()
                    c = self.cache
                    if cr is None:
                        pass
                    elif (c is not None and c['gen'] == P.gen and abs(c['a'] - cr[0]) < 1e-6
                          and abs(c['b'] - cr[1]) < 1e-6):
                        pass                                   # reuse (complete or not)
                    else:
                        self.cache = None
                        a, b = cr
                        if b - a <= CACHE_MAX_S and t0 <= a + fp:
                            w = P.tw if P.zoom is None else min(P.tw, P.zoom[2])
                            h = P.th if P.zoom is None else min(P.th, P.zoom[3])
                            mb = ((b - a) * self.fps + 3) * w * h * 3 / 1e6
                            if mb <= CACHE_MAX_MB:
                                self.cache = dict(a=a, b=b, gen=P.gen, frames=[], times=[],
                                                  complete=False)
                src.seek(t0)
                seek_t = t0
                continue
            if not self.active:
                continue
            with self.cv:
                hit = self._cache_hit(tgt) is not None
                bs = self.buf[0].t if self.buf else None
                be = self.buf[-1].t if self.buf else None
            if hit:
                time.sleep(0.004)
                continue
            dec_t = src.cur[1] if src.cur is not None else seek_t
            back = (bs is not None and tgt < bs - 1.5 * fp) or (bs is None and tgt < dec_t - 1.5 * fp)
            fwd = tgt > max(dec_t, seek_t) + 1.0
            if back or fwd:
                with self.cv:
                    self.buf.clear()
                    c = self.cache
                    if c is not None and not c['complete']:
                        self.cache = None
                src.seek(tgt)
                seek_t = tgt
                continue
            if be is not None and be > tgt + 0.35:
                time.sleep(0.003)
                continue
            x = src.pull()
            if x is None:
                time.sleep(0.01)
                continue
            f, ft = x
            src.cur = x
            c = self.cache
            in_cache = (c is not None and not c['complete'] and c['a'] - fp <= ft <= c['b'] + fp)
            if ft + fp < max(tgt, seek_t) - 1e-4 and not in_cache:
                continue                     # catching up from a keyframe: decode only
            P = self.params.snapshot()
            try:
                p = prepare(f, ft, P)
            except Exception:
                continue
            with self.cv:
                if P.gen != self.params.gen:
                    continue
                self.buf.append(p)
                t_now = self.target
                while len(self.buf) > 1 and self.buf[1].t <= t_now - 0.05:
                    self.buf.popleft()
                c = self.cache
                if c is not None and not c['complete'] and c['gen'] == P.gen:
                    if c['a'] - fp <= ft <= c['b'] + fp:
                        if not c['times'] or ft > c['times'][-1]:
                            c['frames'].append(p)
                            c['times'].append(ft)
                    if ft >= c['b'] - 1e-6 and c['times']:
                        c['complete'] = True
        src.close()


def wsola(x, speed, fs):
    """Slow down / speed up speech WITHOUT changing its pitch (WSOLA)."""
    x = np.asarray(x, dtype=np.float32)
    if abs(speed - 1.0) < 1e-6 or len(x) < int(0.1 * fs):
        return x.copy()
    N = int(0.030 * fs)                      # 30 ms frames
    N -= N % 4
    Hs = N // 2                              # output hop
    Ha = Hs * speed                          # input hop
    tol = int(0.010 * fs)                    # +-10 ms search
    D = 4                                    # search on a 4x decimated copy (fast)
    xp = np.concatenate([x, np.zeros(N + 2 * tol + D, np.float32)])
    xd = xp[:len(xp) // D * D].reshape(-1, D).mean(axis=1)
    win = (0.5 - 0.5 * np.cos(2 * np.pi * np.arange(N) / N)).astype(np.float32)
    nfr = max(1, int((len(x) - N) / Ha) + 1)
    y = np.zeros(nfr * Hs + N, np.float32)
    w = np.zeros_like(y)
    pos = 0
    for k in range(nfr):
        o = k * Hs
        y[o:o + N] += xp[pos:pos + N] * win
        w[o:o + N] += win
        nat = pos + Hs                       # natural continuation of this frame
        nom = int(round((k + 1) * Ha))       # where the next frame should come from
        lo = max(0, nom - tol)
        hi = min(len(x), nom + tol)
        if hi <= lo:
            pos = nom
            continue
        ref = xd[nat // D:(nat + N) // D]
        reg = xd[lo // D:(hi + N) // D]
        if len(reg) < len(ref) or not ref.any():
            pos = nom
            continue
        c = np.correlate(reg, ref, 'valid')
        pos = lo + D * int(np.argmax(c))
    y /= np.maximum(w, 1e-3)
    return y[:int(len(x) / speed)]


FADE_S = 0.004          # short fades: no clicks at start / stop / loop seam
AUDIO_LATENCY = 0.15    # seconds buffered ahead in the audio process
MAX_SLOW_RENDER = 180.0


def render_audio(a, s0, s1, speed, fs):
    seg = a[s0:s1]
    y = wsola(seg, speed, fs) if speed != 1 else np.array(seg, dtype=np.float32)
    f = min(int(FADE_S * fs), len(y) // 4)
    if f > 1:
        ramp = np.linspace(0, 1, f, dtype=np.float32)
        y[:f] *= ramp
        y[-f:] *= ramp[::-1]
    return np.ascontiguousarray(y, dtype=np.float32)


# clock slots shared with the audio process
C_SEQ, C_TDAC, C_POS, C_RATE, C_LEN, C_UNDER, C_ERR, C_READY = range(8)


def _audio_process(conn, clock):
    """Runs in its OWN process: renders and plays audio, nothing else, so
    video decoding and drawing in the main program can never starve it."""
    from multiprocessing import shared_memory
    try:
        import sounddevice as sdd
    except Exception:
        sdd = None
    st = SimpleNamespace(a=None, shm=None, fs=0, B=None, i=0, loop=False, loop_on=True,
                         stream=None, seq=0, rate=0)

    def stop():
        if st.stream is not None:
            try:
                st.stream.abort()
                st.stream.close()
            except Exception:
                pass
            st.stream = None

    def cb(out, frames, ti, status):
        try:
            if status.output_underflow:
                clock[C_UNDER] += 1
        except Exception:
            pass
        B, n, i = st.B, len(st.B), st.i
        try:
            dac, cur = ti.outputBufferDacTime, ti.currentTime
            lag = dac - cur if (dac > 0 and cur > 0 and 0 <= dac - cur < 1) else st.stream.latency
        except Exception:
            lag = AUDIO_LATENCY
        clock[C_TDAC] = time.perf_counter() + lag
        clock[C_POS] = i
        if st.loop and st.loop_on and n > 1:
            j = 0
            while j < frames:
                m = min(frames - j, n - i)
                out[j:j + m, 0] = B[i:i + m]
                j += m
                i = (i + m) % n
            st.i = i
        else:
            m = max(0, min(frames, n - i))
            out[:m, 0] = B[i:i + m]
            out[m:, 0] = 0
            st.i = i + m
            if i + m >= n:
                raise sdd.CallbackStop

    clock[C_READY] = 1.0 if sdd is not None else -1.0
    while True:
        try:
            msg = conn.recv()
        except (EOFError, OSError):
            break
        cmd = msg[0]
        try:
            if cmd == 'load':
                stop()
                _, name, n, fs = msg
                if st.shm is not None:
                    st.shm.close()
                st.shm = shared_memory.SharedMemory(name=name)
                st.a = np.ndarray((n,), np.float32, buffer=st.shm.buf)
                st.fs = fs
                conn.send(('loaded',))
            elif cmd == 'play':
                stop()
                _, seq, s0, s1, speed, loop, i0_src, loop_on = msg
                if sdd is None or st.a is None:
                    clock[C_ERR] = seq
                    continue
                B = render_audio(st.a, s0, s1, speed, st.fs)
                ok = False
                for rate in (st.fs, 48000, 44100):
                    try:
                        st.B = B if rate == st.fs else np.interp(
                            np.arange(int(len(B) * rate / st.fs)) * st.fs / rate, np.arange(len(B)), B
                        ).astype(np.float32)
                        st.rate = rate
                        st.i = int(round(i0_src / speed * rate / st.fs)) % max(1, len(st.B))
                        st.loop, st.loop_on = loop, loop_on
                        clock[C_RATE] = rate
                        clock[C_LEN] = len(st.B)
                        clock[C_POS] = st.i
                        clock[C_TDAC] = time.perf_counter() + AUDIO_LATENCY
                        st.stream = sdd.OutputStream(samplerate=rate, channels=1, dtype='float32',
                                                     callback=cb, latency=AUDIO_LATENCY)
                        st.stream.start()
                        clock[C_SEQ] = seq
                        ok = True
                        break
                    except Exception:
                        st.stream = None
                if not ok:
                    clock[C_ERR] = seq
            elif cmd == 'loop':
                st.loop_on = bool(msg[1])
            elif cmd == 'stop':
                stop()
            elif cmd == 'quit':
                break
        except Exception:
            stop()
    stop()
    if st.shm is not None:
        st.shm.close()


class AudioEngine:
    """Front end for the audio process. Plays [t0, t1] (looping over
    [loopA, t1] if asked) at a speed; slow speeds keep their pitch. Its sample
    clock drives the video. Without an audio device it runs a silent clock."""

    def __init__(self, audio, fs):
        import multiprocessing as mp
        from multiprocessing import shared_memory
        self.fs = fs
        self.n = len(audio)
        self.finished = False
        self.playing = False
        self._loop_on = True
        self.seq = 0
        self.local = True               # silent clock until the audio process confirms
        self.proc = None
        self.shm = None
        self.B_len = 1
        try:
            ctx = mp.get_context('spawn')
            self.clock = ctx.Array('d', 8, lock=False)
            self.conn, child = ctx.Pipe()
            self.proc = ctx.Process(target=_audio_process, args=(child, self.clock), daemon=True)
            self.proc.start()
            self.shm = shared_memory.SharedMemory(create=True, size=max(4, self.n * 4))
            np.ndarray((self.n,), np.float32, buffer=self.shm.buf)[:] = audio
            self.conn.send(('load', self.shm.name, self.n, fs))
            if not self.conn.poll(20):
                raise RuntimeError('audio process did not start')
            self.conn.recv()
            self.has_proc = self.clock[C_READY] > 0
        except Exception:
            self.has_proc = False
            self.close()
        self.silent = not self.has_proc

    @property
    def loop_on(self):
        return self._loop_on

    @loop_on.setter
    def loop_on(self, v):
        self._loop_on = bool(v)
        if self.has_proc:
            try:
                self.conn.send(('loop', self._loop_on))
            except Exception:
                pass

    @property
    def underruns(self):
        return int(self.clock[C_UNDER]) if self.has_proc else 0

    def play(self, t0, t1, speed, loop_a):
        """Start playing; returns the source time at which playback ends."""
        self.stop()
        fs = self.fs
        self.loop = loop_a is not None
        src0 = loop_a if self.loop else t0
        if not self.loop and speed != 1:
            t1 = min(t1, t0 + MAX_SLOW_RENDER)
        s0 = max(0, int(round(src0 * fs)))
        s1 = min(self.n, max(s0 + 2, int(round(t1 * fs))))
        self.src0, self.src1 = s0 / fs, s1 / fs
        self.speed = speed
        self.t_req = t0
        self.finished = False
        self.seq += 1
        self.started_at = time.perf_counter()
        i0_src = int(round((t0 - self.src0) * fs)) if self.loop else 0
        self.local_i0 = i0_src / speed
        self.B_len = max(2, int((s1 - s0) / speed))
        self.local = True
        if self.has_proc:
            self.clock[C_UNDER] = 0
            self.conn.send(('play', self.seq, s0, s1, speed, self.loop, i0_src, self._loop_on))
            self.local = False
        self.silent = not self.has_proc
        self.playing = True
        return self.src1

    def current_time(self):
        if not self.local:
            c = self.clock
            if c[C_ERR] == self.seq:              # no device: fall back to a silent clock
                self.local = True
                self.silent = True
                self.started_at = time.perf_counter()
            elif c[C_SEQ] != self.seq:            # still rendering / starting
                return self.t_req
            else:
                rate, n = c[C_RATE], c[C_LEN]
                o = c[C_POS] + (time.perf_counter() - c[C_TDAC]) * rate
                if self.loop and self._loop_on and n > 1:
                    o %= n
                else:
                    o = max(0.0, min(o, n))
                    if o >= n:
                        self.finished = True
                return min(self.src0 + o / rate * self.speed, self.src1)
        o = self.local_i0 + (time.perf_counter() - self.started_at) * self.fs
        n = self.B_len
        if self.loop and self._loop_on:
            o %= n
        elif o >= n:
            o = n
            self.finished = True
        return min(self.src0 + o / self.fs * self.speed, self.src1)

    def stop(self):
        self.playing = False
        if self.has_proc:
            try:
                self.conn.send(('stop',))
            except Exception:
                pass

    def close(self):
        try:
            if self.proc is not None and self.proc.is_alive():
                self.conn.send(('quit',))
                self.proc.join(1)
                if self.proc.is_alive():
                    self.proc.terminate()
        except Exception:
            pass
        try:
            if self.shm is not None:
                self.shm.close()
                self.shm.unlink()
        except Exception:
            pass
        self.shm = None
        self.proc = None


# ----------------------------------------------------------------------------
# Widgets
# ----------------------------------------------------------------------------
class VideoView(QtWidgets.QWidget):
    """Draws the current frame at full screen resolution with smooth scaling.
    Scroll = zoom, drag = move, double-click = fit."""

    def __init__(self, app):
        super().__init__()
        self.app = app
        self.p = None
        self.msg = 'Click "Load video" to begin'
        self._pan = None
        self.setMinimumSize(240, 140)
        self.setAttribute(Qt.WA_OpaquePaintEvent)
        self.setFocusPolicy(Qt.NoFocus)
        self.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Expanding)

    def set_prepared(self, p):
        self.p = p
        self.update()

    def view_rect(self):
        P = self.app.vparams
        W, H = P.dims()
        z = P.zoom
        return QRectF(0, 0, W, H) if z is None else QRectF(*z)

    def target_rect(self, vr):
        r = self.rect()
        s = min(r.width() / max(vr.width(), 1), r.height() / max(vr.height(), 1))
        w, h = vr.width() * s, vr.height() * s
        return QRectF(r.x() + (r.width() - w) / 2, r.y() + (r.height() - h) / 2, w, h)

    def paintEvent(self, e):
        qp = QPainter(self)
        qp.fillRect(self.rect(), Qt.black)
        if self.p is None or self.app.vparams is None:
            qp.setPen(QColor(150, 150, 150))
            f = qp.font()
            f.setPointSize(14)
            qp.setFont(f)
            qp.drawText(self.rect(), Qt.AlignCenter, self.msg)
            return
        qp.setRenderHint(QPainter.SmoothPixmapTransform, True)
        vr = self.view_rect()
        tr = self.target_rect(vr)
        cx, cy, cw, ch = self.p.cover
        sx = self.p.img.width() / max(cw, 1)
        sy = self.p.img.height() / max(ch, 1)
        src = QRectF((vr.x() - cx) * sx, (vr.y() - cy) * sy, vr.width() * sx, vr.height() * sy)
        qp.drawImage(tr, self.p.img, src)

    def frame_point(self, pos):
        vr = self.view_rect()
        tr = self.target_rect(vr)
        s = vr.width() / max(tr.width(), 1)
        return (vr.x() + (pos.x() - tr.x()) * s, vr.y() + (pos.y() - tr.y()) * s)

    def wheelEvent(self, e):
        if self.app.vparams is None:
            return
        d = e.angleDelta().y()
        if d == 0:
            return
        cx, cy = self.frame_point(e.position())
        self.app.video_zoom_by(1.15 ** (d / 120.0), cx, cy)

    def mouseDoubleClickEvent(self, e):
        if self.app.vparams is not None:
            self.app.reset_video_zoom()

    def mousePressEvent(self, e):
        if self.app.vparams is None:
            return
        if self.app.vparams.zoom is None:
            self.app.set_status('Scroll on the video (or click +) to zoom in; then drag to move around.')
            return
        self._pan = (e.position(), tuple(self.app.vparams.zoom))

    def mouseMoveEvent(self, e):
        if self._pan is None:
            return
        p0, z = self._pan
        vr = QRectF(*z)
        tr = self.target_rect(vr)
        dpp = vr.width() / max(tr.width(), 1)
        d = e.position() - p0
        W, H = self.app.vparams.dims()
        x0 = int(round(min(max(z[0] - d.x() * dpp, 0), W - z[2])))
        y0 = int(round(min(max(z[1] - d.y() * dpp, 0), H - z[3])))
        self.app.set_video_zoom((x0, y0, z[2], z[3]), quiet=True)

    def mouseReleaseEvent(self, e):
        self._pan = None

    def resizeEvent(self, e):
        self.app.schedule_video_params()


class StripBox(pg.ViewBox):
    """ViewBox for the spectrogram / waveform / strips; mouse goes to the app."""

    def __init__(self, app, panel, name):
        super().__init__(enableMenu=False)
        self.app, self.panel, self.name = app, panel, name
        self.setMouseEnabled(False, False)
        self.yr = (0.0, 1.0)

    def _pt(self, sp):
        p = self.mapSceneToView(sp)
        return self.app.clamp_t(p.x()), p.y()

    def mouseClickEvent(self, ev):
        if not self.app.loaded:
            return
        t, y = self._pt(ev.scenePos())
        if ev.button() == Qt.LeftButton:
            ev.accept()
            if ev.double():
                self.app.on_double(self.panel, t, y)
            else:
                self.app.on_click(self.panel, t, y)
        elif ev.button() == Qt.RightButton:
            ev.accept()
            self.app.on_right(self.panel, t, y, ev.screenPos())

    def mouseDragEvent(self, ev, axis=None):
        if ev.button() != Qt.LeftButton or not self.app.loaded:
            ev.ignore()
            return
        ev.accept()
        if ev.isStart():
            t0, y0 = self._pt(ev.buttonDownScenePos())
            self.app.drag_start(self, self.panel, t0, y0, ev.modifiers())
        t, _ = self._pt(ev.scenePos())
        self.app.drag_move(t)
        if ev.isFinish():
            self.app.drag_finish()

    def wheelEvent(self, ev, axis=None):
        ev.accept()
        if self.app.loaded:
            self.app.on_wheel(ev.delta(), ev.modifiers(), self._pt(ev.scenePos())[0])

    def hoverEvent(self, ev):
        if not self.app.loaded or ev.isExit():
            return
        t, y = self._pt(ev.scenePos())
        self.app.on_hover(self, t, y)


class OverviewBox(pg.ViewBox):
    def __init__(self, app):
        super().__init__(enableMenu=False)
        self.app = app
        self.setMouseEnabled(False, False)

    def mouseClickEvent(self, ev):
        if self.app.loaded and ev.button() == Qt.LeftButton:
            ev.accept()
            t = self.app.clamp_t(self.mapSceneToView(ev.scenePos()).x())
            self.app.ov_click(t)

    def mouseDragEvent(self, ev, axis=None):
        if not self.app.loaded or ev.button() != Qt.LeftButton:
            ev.ignore()
            return
        ev.accept()
        if ev.isStart():
            self.app.ov_start(self.app.clamp_t(self.mapSceneToView(ev.buttonDownScenePos()).x()))
        self.app.ov_drag(self.app.clamp_t(self.mapSceneToView(ev.scenePos()).x()))
        if ev.isFinish():
            self.app.ov_end()

    def wheelEvent(self, ev, axis=None):
        ev.accept()


class EventPicker(QtWidgets.QDialog):
    """Tick one or more disfluency / error types. R P B T, 1-6 tick; Enter / Esc."""

    def __init__(self, parent, have, what):
        super().__init__(parent)
        self.setWindowTitle('Disfluency / error')
        self.setModal(True)
        v = QtWidgets.QVBoxLayout(self)
        t = QtWidgets.QLabel('What is %s?' % what)
        f = t.font()
        f.setBold(True)
        f.setPointSize(f.pointSize() + 1)
        t.setFont(f)
        v.addWidget(t)
        v.addWidget(QtWidgets.QLabel('Tick one or more (keys R P B T and 1-6), then Enter.'))
        self.cbx = []
        h = QtWidgets.QLabel('<b>Disfluency</b>')
        v.addWidget(h)
        for q, d in enumerate(DISFL):
            cb = QtWidgets.QCheckBox('%s   %s  -  %s' % (d['key'].upper(), d['abbr'], d['name']))
            cb.setChecked(bool(have[q]))
            cb.setFocusPolicy(Qt.NoFocus)
            c = qc(tuple(x * 0.75 for x in d['color']))
            cb.setStyleSheet('color: %s; font-weight: bold;' % c.name())
            v.addWidget(cb)
            self.cbx.append(cb)
        v.addWidget(QtWidgets.QLabel('<b>Error</b>'))
        nD = len(DISFL)
        for q, e in enumerate(ERRS):
            cb = QtWidgets.QCheckBox('%d   %s  -  %s' % (q + 1, e['abbr'], e['name']))
            cb.setChecked(bool(have[nD + q]))
            cb.setFocusPolicy(Qt.NoFocus)
            cb.setStyleSheet('color: %s;' % qc(tuple(x * 0.8 for x in ERR_COLOR)).name())
            v.addWidget(cb)
            self.cbx.append(cb)
        hb = QtWidgets.QHBoxLayout()
        hb.addStretch()
        ok = QtWidgets.QPushButton('Done (Enter)')
        ca = QtWidgets.QPushButton('Cancel (Esc)')
        for b in (ok, ca):
            b.setFocusPolicy(Qt.NoFocus)
            b.setAutoDefault(False)
        ok.clicked.connect(self.accept)
        ca.clicked.connect(self.reject)
        hb.addWidget(ok)
        hb.addWidget(ca)
        v.addLayout(hb)

    def keyPressEvent(self, e):
        k = e.key()
        if k in (Qt.Key_Return, Qt.Key_Enter):
            self.accept()
            return
        if k == Qt.Key_Escape:
            self.reject()
            return
        ch = e.text().lower()
        for q, d in enumerate(DISFL):
            if ch == d['key']:
                self.cbx[q].toggle()
                return
        if ch.isdigit() and 1 <= int(ch) <= len(ERRS):
            self.cbx[len(DISFL) + int(ch) - 1].toggle()

    def values(self):
        return [cb.isChecked() for cb in self.cbx]


class ErrorPicker(QtWidgets.QDialog):
    def __init__(self, parent, title):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.choice = None
        v = QtWidgets.QVBoxLayout(self)
        v.addWidget(QtWidgets.QLabel('<b>Choose the error type  (keys 1-6, Esc = cancel)</b>'))
        for i, e in enumerate(ERRS):
            b = QtWidgets.QPushButton('%d.   %s   -   %s' % (i + 1, e['abbr'], e['name']))
            b.setFocusPolicy(Qt.NoFocus)
            b.setAutoDefault(False)
            b.clicked.connect(lambda _=False, q=i: self.pick(q))
            v.addWidget(b)
        c = QtWidgets.QPushButton('Cancel  (Esc)')
        c.setFocusPolicy(Qt.NoFocus)
        c.clicked.connect(self.reject)
        v.addWidget(c)

    def pick(self, q):
        self.choice = q
        self.accept()

    def keyPressEvent(self, e):
        if e.key() == Qt.Key_Escape:
            self.reject()
            return
        ch = e.text()
        if ch.isdigit() and 1 <= int(ch) <= len(ERRS):
            self.pick(int(ch) - 1)


# ----------------------------------------------------------------------------
# Main window
# ----------------------------------------------------------------------------
class MainWindow(QtWidgets.QMainWindow):

    # ======================= setup ===========================================
    def __init__(self, init_path=None):
        super().__init__()
        self.settings = QSettings('DisfluencyAnnotator', 'DisfluencyAnnotator')
        self.loaded = False
        self.video_path = ''
        self.audio = None
        self.fs = 0
        self.dur = 1.0
        self.amax = 1.0
        self.max_freq = 5000
        self.viewS, self.viewE = 0.0, 1.0
        self.cursor = 0.0
        self.sel_s, self.sel_e = NAN, NAN
        self.events = []
        self.next_uid = 1
        self.cur = None
        self.dirty = False
        self.undo_stack = []
        self.pending_new = None
        self.group_u = 0
        self.expanded = set()
        self.unresolved_open = False
        self.lane_of = []
        self.nlane_p = 1
        self.nlane_e = 1
        self.editing = False
        self.box_target = None
        self.box_kind = ''
        self.drag = None
        self.ov_mode = ''
        self.ov_grab = 0.0
        self.playing = False
        self.play_end = 0.0
        self.play_loop_a = None
        self.play_speed = 1.0
        self.loop_on = True
        self.last_shown = None
        self.ev_items = []
        self.ov_items = []
        try:
            self.contrast = min(80, max(0, int(self.settings.value('specContrast', 30))))
        except Exception:
            self.contrast = 30
        self.spec_max = NAN
        self.vsrc = None
        self.worker = None
        self.aeng = None
        self.vparams = None
        self.last_raw = None
        self.last_raw_t = 0.0
        self.pop = None

        self._build_ui()
        self.play_timer = QTimer(self)
        self.play_timer.setTimerType(Qt.PreciseTimer)
        self.play_timer.setInterval(8)
        self.play_timer.timeout.connect(lambda: self.cb(self.play_tick))
        self.vp_timer = QTimer(self)
        self.vp_timer.setSingleShot(True)
        self.vp_timer.setInterval(120)
        self.vp_timer.timeout.connect(lambda: self.cb(self.apply_video_params))
        QtWidgets.QApplication.instance().installEventFilter(self)
        self.set_video_dependent(False)
        self.show()
        QTimer.singleShot(50, self.show_welcome)
        if init_path:
            QTimer.singleShot(200, lambda: self.cb(lambda: self.load_video(init_path)))

    def cb(self, fn):
        """Run a callback; show errors in a dialog instead of crashing."""
        try:
            fn()
        except Exception as err:
            import traceback
            tb = traceback.format_exc()
            sys.stderr.write(tb)
            self.set_status('ERROR: %s' % err, 'warn')
            QtWidgets.QMessageBox.critical(self, 'Error', '%s\n\n%s' % (err, tb.splitlines()[-3]))

    def _btn(self, label, fn, tip='', parent_layout=None):
        b = QtWidgets.QPushButton(label)
        b.setFocusPolicy(Qt.NoFocus)
        b.setToolTip(tip)
        b.clicked.connect(lambda: self.cb(lambda: self._btn_press(fn)))
        if parent_layout is not None:
            parent_layout.addWidget(b)
        return b

    def _btn_press(self, fn):
        if self.editing:
            self.commit_edit()
        fn()

    def _build_ui(self):
        self.setWindowTitle(APPNAME)
        scr = QtWidgets.QApplication.primaryScreen().availableGeometry()
        self.resize(min(1400, scr.width() - 60), min(960, scr.height() - 80))
        cw = QtWidgets.QWidget()
        cw.setFocusPolicy(Qt.StrongFocus)
        self.setCentralWidget(cw)
        outer = QtWidgets.QHBoxLayout(cw)
        outer.setContentsMargins(8, 8, 8, 4)
        left = QtWidgets.QVBoxLayout()
        right = QtWidgets.QVBoxLayout()
        outer.addLayout(left, 78)
        outer.addLayout(right, 22)

        # ---- top row: annotation panel + IPA | video ----
        top = QtWidgets.QHBoxLayout()
        left.addLayout(top, 40)
        annCol = QtWidgets.QVBoxLayout()
        top.addLayout(annCol, 30)
        gb = QtWidgets.QGroupBox('Selected annotation')
        gv = QtWidgets.QVBoxLayout(gb)
        self.annot = QtWidgets.QPlainTextEdit()
        self.annot.setReadOnly(True)
        self.annot.setFocusPolicy(Qt.NoFocus)
        f = QFont('Menlo' if sys.platform == 'darwin' else 'Consolas')
        f.setStyleHint(QFont.Monospace)
        f.setPointSize(11)
        self.annot.setFont(f)
        gv.addWidget(self.annot)
        annCol.addWidget(gb, 1)
        ipa = QtWidgets.QGridLayout()
        ipa.setSpacing(2)
        for r, syms in enumerate(IPA_ROWS):
            for c, sym in enumerate(syms):
                b = QtWidgets.QPushButton(sym)
                b.setFocusPolicy(Qt.NoFocus)
                b.setToolTip('Insert into the text you are typing')
                b.setMinimumWidth(24)
                fb = b.font()
                fb.setPointSize(13)
                b.setFont(fb)
                b.clicked.connect(lambda _=False, s=sym: self.cb(lambda: self.insert_symbol(s)))
                ipa.addWidget(b, r, c)
        annCol.addLayout(ipa)

        self.vid_col = QtWidgets.QVBoxLayout()
        top.addLayout(self.vid_col, 62)
        vbar = QtWidgets.QHBoxLayout()
        vbar.addStretch()
        self.btn_pop = self._btn('Pop out video', self.toggle_pop,
                                 'Open the video in its own window (resize it for a bigger picture)', vbar)
        self.btn_vin = self._btn('+', lambda: self.video_zoom_by(1.5), 'Zoom the video in (or scroll on it)', vbar)
        self.btn_vout = self._btn('-', lambda: self.video_zoom_by(1 / 1.5), 'Zoom the video out', vbar)
        self.btn_vfit = self._btn('Fit', self.reset_video_zoom, 'Show the whole picture (or double-click it)', vbar)
        self.btn_rot = self._btn('Rotate', self.rotate_video, 'Rotate the picture 90 degrees (remembered per video)', vbar)
        self.vid_bar = vbar
        self.vid_col.addLayout(vbar)
        self.video = VideoView(self)
        self.vid_col.addWidget(self.video, 1)
        self.dock_msg = QtWidgets.QLabel('Video is in its own window.\n\nClick "Dock video" '
                                         '(or close that window) to bring it back here.')
        self.dock_msg.setAlignment(Qt.AlignCenter)
        self.dock_msg.setStyleSheet('background:#262626;color:white;font-size:13px;')
        self.dock_msg.hide()
        self.vid_col.addWidget(self.dock_msg, 1)

        # ---- timeline ----
        self.glw = pg.GraphicsLayoutWidget()
        self.glw.setBackground(qc(BG))
        self.glw.setFocusPolicy(Qt.StrongFocus)
        left.addWidget(self.glw, 46)
        specs = [('spec', '', 'Freq (Hz)', 50), ('wave', '', 'Amp', 25),
                 ('trans', 'transcript', 'Transcript', 18), ('de', 'disferr', 'Disfl/Error', 28),
                 ('notes', 'notes', 'Notes', 22)]
        self.vbs = {}
        for i, (name, panel, label, stretch) in enumerate(specs):
            vb = StripBox(self, panel, name)
            pi = pg.PlotItem(viewBox=vb)
            pi.hideButtons()
            pi.setMenuEnabled(False)
            pi.setLabel('left', label)
            ax = pi.getAxis('left')
            ax.setWidth(72)
            if name != 'spec':
                ax.setStyle(showValues=False, tickLength=0)
            if name != 'notes':
                pi.hideAxis('bottom')
            else:
                pi.setLabel('bottom', 'Time (s)')
            vb.setBackgroundColor('w')
            vb.setBorder(pg.mkPen((40, 40, 40)))
            vb.disableAutoRange()
            vb.setDefaultPadding(0)
            self.glw.addItem(pi, row=i, col=0)
            self.glw.ci.layout.setRowStretchFactor(i, stretch)
            self.vbs[name] = vb
            if i > 0:
                vb.setXLink(self.vbs['spec'])
        self.glw.ci.layout.setVerticalSpacing(2)
        self.vb_spec, self.vb_wave = self.vbs['spec'], self.vbs['wave']
        self.vb_trans, self.vb_de, self.vb_notes = self.vbs['trans'], self.vbs['de'], self.vbs['notes']
        for name, yr in (('wave', (-1, 1)), ('trans', (0, 1)), ('de', (0, 1)), ('notes', (0, 1)),
                         ('spec', (0, 5000))):
            self.vbs[name].yr = yr
            self.vbs[name].setYRange(*yr, padding=0)
        self.spec_img = pg.ImageItem()
        self.spec_img.setZValue(-10)
        lut = np.repeat(np.arange(255, -1, -1, dtype=np.uint8)[:, None], 3, axis=1)
        self.spec_img.setLookupTable(lut)
        self.vb_spec.addItem(self.spec_img)
        self.wave_curve = pg.PlotCurveItem(pen=pg.mkPen((20, 20, 20), width=1))
        self.wave_curve.setZValue(-10)
        self.vb_wave.addItem(self.wave_curve)
        self.sel_items = []
        for vb in self.vbs.values():
            r = QtWidgets.QGraphicsRectItem()
            r.setBrush(qc(SEL_COLOR, 0.12))
            r.setPen(QPen(Qt.NoPen))
            r.setZValue(5)
            vb.addItem(r, ignoreBounds=True)
            ls = []
            for _ in range(2):
                ln = pg.InfiniteLine(angle=90, pen=pg.mkPen('k', width=2.5))
                ln.setZValue(6)
                vb.addItem(ln, ignoreBounds=True)
                ls.append(ln)
            r.hide()
            ls[0].hide()
            ls[1].hide()
            self.sel_items.append((vb, r, ls[0], ls[1]))
        self.cur_lines = []
        self.play_lines = []
        for vb in (self.vb_spec, self.vb_wave):
            c = pg.InfiniteLine(angle=90, pen=pg.mkPen(qc(CUR_COLOR), width=1, style=Qt.DashLine))
            c.setZValue(8)
            vb.addItem(c, ignoreBounds=True)
            self.cur_lines.append(c)
            p = pg.InfiniteLine(angle=90, pen=pg.mkPen(qc(PLAY_COLOR), width=1.5))
            p.setZValue(9)
            p.hide()
            vb.addItem(p, ignoreBounds=True)
            self.play_lines.append(p)
        # inline typing box
        self.box = QtWidgets.QLineEdit(self.glw.viewport())
        self.box.hide()
        self.box.setStyleSheet('background:white;color:black;border:2px solid #333;font-size:13px;')
        self.box.returnPressed.connect(lambda: self.cb(self.commit_box))

        # ---- view row ----
        row = QtWidgets.QHBoxLayout()
        self.view_lbl = QtWidgets.QLabel('')
        row.addWidget(self.view_lbl, 1)
        row.addWidget(QtWidgets.QLabel('Show from'))
        self.edt_from = QtWidgets.QLineEdit()
        self.edt_to = QtWidgets.QLineEdit()
        for e, tag in ((self.edt_from, 'from'), (self.edt_to, 'to')):
            e.setFixedWidth(90)
            e.setAlignment(Qt.AlignCenter)
            e.setToolTip('Time (seconds, m:ss or 2m30) - type, Enter')
            e.returnPressed.connect(lambda t=tag: self.cb(lambda: self.apply_time_fields(t)))
        row.addWidget(self.edt_from)
        row.addWidget(QtWidgets.QLabel('to'))
        row.addWidget(self.edt_to)
        self.btn_go = self._btn('Go', lambda: self.apply_time_fields(''),
                                'Show the time frame typed in the From / To boxes', row)
        left.addLayout(row)

        # ---- overview ----
        self.ov_vb = OverviewBox(self)
        self.ovw = pg.PlotWidget(viewBox=self.ov_vb)
        self.ovw.setBackground(qc(BG))
        self.ovw.setFixedHeight(52)
        self.ovw.hideButtons()
        self.ovw.getPlotItem().setMenuEnabled(False)
        self.ovw.getPlotItem().hideAxis('bottom')
        self.ovw.getPlotItem().setLabel('left', 'Overview')
        self.ovw.getAxis('left').setWidth(72)
        self.ovw.getAxis('left').setStyle(showValues=False, tickLength=0)
        self.ov_vb.setBackgroundColor('w')
        self.ov_vb.setBorder(pg.mkPen((40, 40, 40)))
        self.ov_vb.disableAutoRange()
        self.ov_vb.setYRange(0, 1, padding=0)
        self.ov_env = pg.PlotCurveItem(pen=pg.mkPen((140, 140, 140), width=1))
        self.ov_vb.addItem(self.ov_env)
        self.ov_rect = QtWidgets.QGraphicsRectItem()
        self.ov_rect.setBrush(qc(OV_COLOR, 0.25))
        pen = QPen(qc(OV_COLOR))
        pen.setCosmetic(True)
        pen.setWidthF(1.5)
        self.ov_rect.setPen(pen)
        self.ov_vb.addItem(self.ov_rect, ignoreBounds=True)
        self.ov_cur = pg.InfiniteLine(angle=90, pen=pg.mkPen(qc(CUR_COLOR), width=1.2))
        self.ov_vb.addItem(self.ov_cur, ignoreBounds=True)
        left.addWidget(self.ovw)

        # ---- navigation row ----
        nav = QtWidgets.QHBoxLayout()
        lbl = QtWidgets.QLabel('<b>Navigate:</b>')
        nav.addWidget(lbl)
        self.btn_zin = self._btn('Zoom in (Cmd +)', lambda: self.zoom_about(self.zoom_center(), 0.5),
                                 'Zoom in on the highlight / cursor', nav)
        self.btn_zout = self._btn('Zoom out (Cmd -)', lambda: self.zoom_about(self.zoom_center(), 2),
                                  'Zoom out', nav)
        self.btn_zsel = self._btn('Fit selection (Z)', self.zoom_to_selection, 'Zoom to the highlight (Z)', nav)
        self.btn_zall = self._btn('Full recording (A)', self.full_view, 'Show the whole recording (A)', nav)
        self.chk_loop = QtWidgets.QCheckBox('Loop (L)')
        self.chk_loop.setChecked(True)
        self.chk_loop.setFocusPolicy(Qt.NoFocus)
        self.chk_loop.toggled.connect(lambda v: self.cb(lambda: self.set_loop(v, quiet=False)))
        nav.addWidget(self.chk_loop)
        self.btn_play = self._btn('Play (Space)', self.toggle_play, 'Play / pause the highlight (Space)', nav)
        nav.addStretch()
        nav.addWidget(QtWidgets.QLabel('Contrast'))
        self.sld = QtWidgets.QSlider(Qt.Horizontal)
        self.sld.setRange(0, 80)
        self.sld.setValue(self.contrast)
        self.sld.setFixedWidth(130)
        self.sld.setFocusPolicy(Qt.NoFocus)
        self.sld.setToolTip('Spectrogram contrast - right = darker / more contrast. Keys [ and ]')
        self.sld.valueChanged.connect(lambda v: self.cb(lambda: self.set_contrast(v)))
        nav.addWidget(self.sld)
        left.addLayout(nav)

        # ---- file row + legend ----
        frow = QtWidgets.QHBoxLayout()
        self._btn('Load video', self.ui_load_video, 'Open a video file', frow)
        self.btn_load_ann = self._btn('Load annotations...', self.load_annotations,
                                      'Load an annotation table (checks it first); merge or replace', frow)
        self.btn_save = self._btn('Save annotations', self.save_annotations,
                                  'Save all annotations to an Excel file (Ctrl+S)', frow)
        self.btn_undo = self._btn('Undo', self.undo, 'Undo the last change (Ctrl+Z)', frow)
        self._btn('Help', self.show_help, 'Full instructions (F1)', frow)
        self._btn('Shortcuts', self.show_shortcuts, 'Keyboard and mouse shortcuts', frow)
        frow.addStretch()
        legs = [('%s fluent' % FLU_KEY.upper(), FLUENT_COLOR)] + \
               [('%s %s' % (d['key'].upper(), d['abbr']), d['color']) for d in DISFL] + \
               [('%s error' % ERR_KEY.upper(), ERR_COLOR)]
        for name, col in legs:
            sw = QtWidgets.QLabel()
            sw.setFixedSize(12, 12)
            sw.setStyleSheet('background:%s;' % qc(col).name())
            frow.addWidget(sw)
            frow.addWidget(QtWidgets.QLabel(name + '  '))
        left.addLayout(frow)

        # ---- right column ----
        ev_gb = QtWidgets.QGroupBox('Fluent events (click to open)')
        ev_l = QtWidgets.QVBoxLayout(ev_gb)
        self.tree = QtWidgets.QTreeWidget()
        self.tree.setHeaderLabels(['Start', 'End', 'Fluent event', 'Notes'])
        self.tree.setFocusPolicy(Qt.NoFocus)
        self.tree.setRootIsDecorated(True)
        self.tree.setIndentation(12)
        self.tree.setColumnWidth(0, 92)
        self.tree.setColumnWidth(1, 76)
        self.tree.setColumnWidth(2, 100)
        self.tree.itemClicked.connect(self.on_tree_click)
        self.tree.itemExpanded.connect(lambda it: self.on_tree_expand(it, True))
        self.tree.itemCollapsed.connect(lambda it: self.on_tree_expand(it, False))
        ev_l.addWidget(self.tree, 1)
        g = QtWidgets.QGridLayout()
        self.btn_play_ev = self._btn('Play event', self.play_current_event, 'Play the selected item')
        self.btn_edit_ev = self._btn('Edit text', self.edit_current_event, 'Type the text of the selected item')
        self.btn_type_ev = self._btn('Change type', self.change_type_dialog, 'Change the type of the selected part / error')
        self.btn_del_ev = self._btn('Delete event', self.delete_current_event, 'Delete the selected item (Ctrl+Z to undo)')
        self.btn_add_txt = self._btn('Transcript', lambda: self.begin_new_text(),
                                     'Type the transcript for the selection (or double-click the highlight)')
        self.btn_add_flu = self._btn('Fluent part (F)', lambda: self.on_part_key('flu'),
                                     'Mark the selection as the fluent word (inside a transcript)')
        for i, b in enumerate([self.btn_play_ev, self.btn_edit_ev, self.btn_type_ev, self.btn_del_ev,
                               self.btn_add_txt, self.btn_add_flu]):
            g.addWidget(b, i // 2, i % 2)
        ev_l.addLayout(g)
        sp = QtWidgets.QHBoxLayout()
        sp.addWidget(QtWidgets.QLabel('Playback speed:'))
        self.pop_speed = QtWidgets.QComboBox()
        self.pop_speed.addItems(SPEED_NAMES)
        self.pop_speed.setFocusPolicy(Qt.NoFocus)
        self.pop_speed.activated.connect(lambda i: self.cb(lambda: self.set_speed_idx(i + 1)))
        sp.addWidget(self.pop_speed)
        ev_l.addLayout(sp)
        self.info = QtWidgets.QLabel('No selection')
        self.info.setWordWrap(True)
        ev_l.addWidget(self.info)
        right.addWidget(ev_gb, 3)
        tips = QtWidgets.QGroupBox('Quick tips')
        tl = QtWidgets.QVBoxLayout(tips)
        tlab = QtWidgets.QLabel('\n'.join([
            'New speech: highlight + double-click = transcript.',
            'Notes: double-click highlight on Notes strip.',
            'Word: highlight + Enter = fluent event selected.',
            'Part: highlight + Enter = tick type(s), Enter.',
            'W = whole episode. Esc = done. N = notes.',
            'Zoom: Cmd/Ctrl + / -.  Z = highlight, A = all.',
            'Contrast: [ less, ] more (or the slider).',
            'Space plays (L loop, 1-4 speed). F1 = help']))
        tlab.setWordWrap(True)
        f = tlab.font()
        f.setPointSize(max(8, f.pointSize() - 1))
        tlab.setFont(f)
        tl.addWidget(tlab)
        right.addWidget(tips, 1)

        self.status = QtWidgets.QLabel('No video loaded - click "Load video" to begin.')
        self.statusBar().addWidget(self.status, 1)

        self.video_dependent = [self.edt_from, self.edt_to, self.btn_go, self.tree, self.btn_play_ev,
                                self.btn_edit_ev, self.btn_type_ev, self.btn_del_ev, self.btn_add_txt,
                                self.btn_add_flu, self.pop_speed, self.btn_zin, self.btn_zout,
                                self.btn_zsel, self.btn_zall, self.chk_loop, self.btn_play, self.sld,
                                self.btn_load_ann, self.btn_save, self.btn_undo, self.btn_pop,
                                self.btn_vin, self.btn_vout, self.btn_vfit, self.btn_rot]

    def set_video_dependent(self, on):
        for w in self.video_dependent:
            w.setEnabled(on)

    # ======================= keyboard =========================================
    def eventFilter(self, obj, e):
        if e.type() != QEvent.KeyPress:
            return False
        aw = QtWidgets.QApplication.activeWindow()
        if aw is not self and (self.pop is None or aw is not self.pop):
            return False
        fw = QtWidgets.QApplication.focusWidget()
        if isinstance(fw, QtWidgets.QLineEdit):
            if e.key() == Qt.Key_Escape:
                if fw is self.box:
                    self.cb(self.cancel_typing)
                    return True
                if fw in (self.edt_from, self.edt_to):
                    self.sync_time_fields(self.viewS, self.viewE, force=True)
                    self.glw.setFocus()
                    return True
            return False
        if e.isAutoRepeat() and e.key() not in (Qt.Key_Left, Qt.Key_Right, Qt.Key_BracketLeft,
                                                Qt.Key_BracketRight):
            return True
        handled = [False]

        def run():
            handled[0] = self.handle_key(e)
        self.cb(run)
        return handled[0]

    def handle_key(self, e):
        k = e.key()
        mods = e.modifiers()
        if k == Qt.Key_F1:
            self.show_help()
            return True
        letters = {Qt.Key_Space, Qt.Key_R, Qt.Key_B, Qt.Key_P, Qt.Key_T, Qt.Key_E, Qt.Key_F,
                   Qt.Key_M, Qt.Key_W, Qt.Key_N, Qt.Key_Return, Qt.Key_Enter}
        if not self.loaded:
            if k in letters:
                self.set_status('Load a video first (click "Load video").', 'warn')
                return True
            return False
        ctrl = bool(mods & (Qt.ControlModifier | Qt.MetaModifier))
        shift = bool(mods & Qt.ShiftModifier)
        if ctrl:
            if k in (Qt.Key_Equal, Qt.Key_Plus):
                self.zoom_about(self.zoom_center(), 0.5)
            elif k in (Qt.Key_Minus, Qt.Key_Underscore):
                self.zoom_about(self.zoom_center(), 2)
            elif k == Qt.Key_O:
                self.zoom_about(self.cursor_center(), 2)
            elif k == Qt.Key_I:
                self.zoom_about(self.cursor_center(), 0.5)
            elif k == Qt.Key_N:
                self.zoom_to_selection()
            elif k == Qt.Key_A:
                self.full_view()
            elif k == Qt.Key_S:
                self.save_annotations()
            elif k == Qt.Key_Z:
                self.undo()
            else:
                return False
            return True
        ch = e.text().lower()
        if k == Qt.Key_Space:
            self.toggle_play()
        elif k == Qt.Key_Escape:
            if self.playing:
                self.stop_play()
                self.set_status('Stopped.')
            elif self.group_u > 0:
                self.exit_group()
            else:
                self.clear_selection()
        elif k in (Qt.Key_Return, Qt.Key_Enter):
            self.on_enter()
        elif k in (Qt.Key_Delete, Qt.Key_Backspace):
            self.delete_current_event()
        elif k == Qt.Key_Left:
            self.step_cursor(-1, shift)
        elif k == Qt.Key_Right:
            self.step_cursor(1, shift)
        elif k == Qt.Key_BracketLeft:
            self.set_contrast(self.contrast - 5)
        elif k == Qt.Key_BracketRight:
            self.set_contrast(self.contrast + 5)
        elif ch == ERR_KEY:
            self.start_add_error()
        elif ch == FLU_KEY:
            self.on_part_key('flu')
        elif ch == 'm':
            self.open_event_picker()
        elif ch == 'w':
            self.whole_episode()
        elif ch == 'n':
            self.edit_notes()
        elif ch == 'l':
            self.set_loop(not self.loop_on)
        elif ch in ('1', '2', '3', '4'):
            self.set_speed_idx(int(ch))
        elif ch == 'z':
            self.zoom_to_selection()
        elif ch == 'a':
            self.full_view()
        else:
            for d in DISFL:
                if ch == d['key']:
                    self.on_part_key(d['abbr'])
                    return True
            return False
        return True

    # ======================= model helpers ====================================
    def find_uid(self, u):
        if not u or u <= 0:
            return None
        for i, e in enumerate(self.events):
            if e.uid == u:
                return i
        return None

    def new_event(self, a, b, kind, cat, tr, link):
        self.events.append(Ev(a, b, kind, cat, tr, '', self.next_uid, link))
        self.next_uid += 1
        return len(self.events) - 1

    def overlap_idx(self, a, b, kinds, exclude=()):
        return [i for i, e in enumerate(self.events)
                if e.start < b - TOL and e.end > a + TOL and e.kind in kinds and i not in exclude]

    def seg_containing(self, a, b):
        for i, e in enumerate(self.events):
            if e.kind == 'trans' and e.start <= a + TOL and e.end >= b - TOL:
                return i
        return None

    def seg_of(self, k):
        if k is None or k >= len(self.events) or self.events[k].kind == 'trans':
            return None
        return self.seg_containing(self.events[k].start, self.events[k].end)

    def fluents_in_seg(self, sk):
        if sk is None:
            return []
        s = self.events[sk]
        idx = [i for i, e in enumerate(self.events)
               if e.kind == 'fluent' and e.start >= s.start - TOL and e.end <= s.end + TOL]
        return sorted(idx, key=lambda i: self.events[i].start)

    def linked_all(self, u):
        if u <= 0:
            return []
        return [i for i, e in enumerate(self.events) if e.kind in ('disfl', 'error') and e.link == u]

    def group_uid(self, k):
        if k is None or k >= len(self.events):
            return 0
        e = self.events[k]
        if e.kind == 'fluent':
            return e.uid
        if e.kind in ('disfl', 'error'):
            if e.link > 0 and self.find_uid(e.link) is not None:
                return e.link
            return 0
        for j in self.attached_to(k):
            u = self.group_uid(j)
            if u > 0:
                return u
        return 0

    def attached_to(self, k):
        if k is None or k >= len(self.events) or self.events[k].kind != 'trans':
            return []
        u = self.events[k].uid
        return [i for i, e in enumerate(self.events) if e.tok == u and e.kind != 'trans']

    def sync_tok(self, k):
        for j in self.attached_to(k):
            self.events[j].start = self.events[k].start
            self.events[j].end = self.events[k].end
            self.clear_global(j)

    def clear_global(self, k):
        self.events[k].gStart = NAN
        self.events[k].gStop = NAN

    def link_for(self, k):
        fl = [i for i, e in enumerate(self.events) if e.kind == 'fluent']
        if not fl:
            return 0
        a, b = self.events[k].start, self.events[k].end
        ov = [min(b, self.events[i].end) - max(a, self.events[i].start) for i in fl]
        if max(ov) > TOL:
            return self.events[fl[int(np.argmax(ov))]].uid
        after = [i for i in fl if self.events[i].start >= b - TOL and self.events[i].start - b <= 1.5]
        if after:
            return self.events[min(after, key=lambda i: self.events[i].start)].uid
        return 0

    def relink_all(self):
        for j, e in enumerate(self.events):
            if e.kind in ('disfl', 'error'):
                if e.link > 0 and self.find_uid(e.link) is not None:
                    continue
                e.link = self.link_for(j)

    def group_span(self, u):
        m = [self.find_uid(u)] + self.linked_all(u)
        return min(self.events[i].start for i in m), max(self.events[i].end for i in m)

    def partition_ok(self, k):
        if k is None or k >= len(self.events):
            return True
        e = self.events[k]
        if e.kind == 'trans':
            return not self.overlap_idx(e.start, e.end, ['trans'], [k])
        if e.kind == 'fluent':
            return not self.overlap_idx(e.start, e.end, ['fluent'], [k])
        return True

    def has_annot(self, a, b, kind, cat):
        for i, e in enumerate(self.events):
            if e.kind == kind and e.cat == cat and abs(e.start - a) < 1e-3 and abs(e.end - b) < 1e-3:
                return i
        return None

    def piece_types(self, a, b):
        labs = [('fluent' if e.kind == 'fluent' else e.cat) for e in self.events
                if e.kind != 'trans' and abs(e.start - a) < 1e-3 and abs(e.end - b) < 1e-3]
        return ' + '.join(labs) if labs else 'nothing'

    def event_color(self, k):
        e = self.events[k]
        if e.kind == 'fluent':
            return FLUENT_COLOR
        if e.kind == 'error':
            return ERR_COLOR
        if e.kind == 'trans':
            return TRANS_COLOR
        for d in DISFL:
            if d['abbr'] == e.cat:
                return d['color']
        return (0.5, 0.5, 0.5)

    def piece_colors(self, k):
        pc, keys = [], []
        own = self.attached_to(k)
        a, b = self.events[k].start, self.events[k].end
        over = [i for i, e in enumerate(self.events) if e.kind in ('disfl', 'error') and e.tok == 0
                and (min(b, e.end) - max(a, e.start)) >= 0.5 * (b - a)]
        m = own + over
        for f in own:
            if self.events[f].kind == 'fluent':
                ln = sorted(self.linked_all(self.events[f].uid), key=lambda i: self.events[i].start)
                m += ln
        for j in m:
            e = self.events[j]
            if e.kind == 'fluent':
                continue
            key = 'error' if e.kind == 'error' else e.kind + e.cat
            if key in keys:
                continue
            keys.append(key)
            pc.append(self.event_color(j))
        return pc

    def compute_lanes(self):
        n = len(self.events)
        self.lane_of = [0] * n
        P = [i for i in range(n) if self.events[i].kind == 'disfl']
        l, c = self._assign_lanes(P)
        for i, v in zip(P, l):
            self.lane_of[i] = v
        self.nlane_p = max(1, c)
        E = [i for i in range(n) if self.events[i].kind == 'error']
        l, c = self._assign_lanes(E)
        for i, v in zip(E, l):
            self.lane_of[i] = v
        self.nlane_e = max(1, c)

    def _assign_lanes(self, idx):
        lane = [0] * len(idx)
        if not idx:
            return lane, 0
        order = sorted(range(len(idx)), key=lambda j: (self.events[idx[j]].start,
                                                       -(self.events[idx[j]].end - self.events[idx[j]].start)))
        ends = []
        for j in order:
            e = self.events[idx[j]]
            L = next((q for q, en in enumerate(ends) if en <= e.start + TOL), None)
            if L is None:
                ends.append(e.end)
                L = len(ends) - 1
            else:
                ends[L] = e.end
            lane[j] = L + 1
        return lane, len(ends)

    def lane_band(self, k):
        if len(self.lane_of) != len(self.events):
            self.compute_lanes()
        if self.events[k].kind == 'error':
            r, n = (0.02, 0.34), self.nlane_e
        else:
            r, n = (0.36, 0.98), self.nlane_p
        L = max(1, self.lane_of[k])
        h = (r[1] - r[0]) / n
        top = r[1] - (L - 1) * h
        return (top - h + 0.008, top - 0.008)

    def event_at_time(self, t, panel='', y=NAN):
        hit = [i for i, e in enumerate(self.events) if e.start <= t <= e.end]
        if not hit:
            return None
        kind = lambda i: self.events[i].kind
        if panel == 'transcript':
            pref = [i for i in hit if kind(i) == 'trans']
        elif panel == 'disferr':
            cand = [i for i in hit if kind(i) not in ('trans', 'fluent')]
            pref = []
            if not isnan(y):
                for c in cand:
                    b = self.lane_band(c)
                    if b[0] - 0.01 <= y <= b[1] + 0.01:
                        pref.append(c)
            if not pref:
                pref = cand
        elif panel == 'notes':
            pref = hit
        else:
            pref = [i for i in hit if kind(i) not in ('trans', 'fluent')]
            if not pref:
                pref = [i for i in hit if kind(i) == 'trans']
        if not pref:
            return None
        return min(pref, key=lambda i: self.events[i].end - self.events[i].start)

    def piece_under(self, a, b, t):
        tr = [i for i, e in enumerate(self.events) if e.kind == 'trans']
        if not tr:
            return None
        at = [i for i in tr if self.events[i].start <= t <= self.events[i].end]
        if at:
            return at[0]
        ov = [min(b, self.events[i].end) - max(a, self.events[i].start) for i in tr]
        m = int(np.argmax(ov))
        return tr[m] if ov[m] > TOL else None

    def touching_piece(self, k, side):
        for j, e in enumerate(self.events):
            if j == k or e.kind != 'trans':
                continue
            if side == 'left' and abs(e.end - self.events[k].start) < 1e-3:
                return j
            if side == 'right' and abs(e.start - self.events[k].end) < 1e-3:
                return j
        return None

    def focus_set(self):
        if not self.events:
            return []
        if self.group_u > 0 and self.find_uid(self.group_u) is not None:
            u = self.group_u
        else:
            if self.cur is None or self.cur >= len(self.events):
                return []
            u = self.group_uid(self.cur)
        if u <= 0:
            return []
        m = [self.find_uid(u)] + self.linked_all(u)
        for tku in {self.events[i].tok for i in m if self.events[i].tok > 0}:
            j = self.find_uid(tku)
            if j is not None:
                m.append(j)
        return [i for i in set(m) if i != self.cur]

    def words_under(self, a, b):
        tr = [i for i, e in enumerate(self.events) if e.kind == 'trans']
        if not tr:
            return ''
        ov = [min(b, self.events[i].end) - max(a, self.events[i].start) for i in tr]
        mx = max(ov)
        use = [i for i, o in zip(tr, ov) if o > 0 and
               (o > 0.3 * (self.events[i].end - self.events[i].start) or o == mx)]
        use.sort(key=lambda i: self.events[i].start)
        return ' '.join(token_word(self.events[i].transcript) for i in use)

    def add_pieces(self, a, b, s):
        toks = split_tokens(s)
        w = np.maximum(1, [len(t) for t in toks]).astype(float)
        edges = a + (b - a) * np.concatenate([[0], np.cumsum(w)]) / w.sum()
        k1 = None
        for i, t in enumerate(toks):
            k = self.new_event(float(edges[i]), float(edges[i + 1]), 'trans', '', t, 0)
            if k1 is None:
                k1 = k
        return k1

    def add_raw(self, a, b, tk, kind, cat):
        txt = ''
        if tk is not None and kind != 'error':
            txt = token_word(self.events[tk].transcript)
        lk = self.group_u if (self.group_u > 0 and kind in ('disfl', 'error')) else 0
        k = self.new_event(a, b, kind, cat, txt, lk)
        if tk is not None:
            self.events[k].tok = self.events[tk].uid
        return k

    def has_sel(self):
        return not isnan(self.sel_s) and not isnan(self.sel_e) and self.sel_e > self.sel_s

    def current_target(self):
        if self.cur is not None and self.cur < len(self.events):
            e = self.events[self.cur]
            if e.kind == 'trans':
                return e.start, e.end, self.cur
            if e.tok > 0:
                t2 = self.find_uid(e.tok)
                if t2 is not None:
                    return self.events[t2].start, self.events[t2].end, t2
            return e.start, e.end, None
        if self.has_sel():
            a, b = self.sel_s, self.sel_e
            tk = next((i for i, e in enumerate(self.events) if e.kind == 'trans'
                       and abs(e.start - a) < 1e-3 and abs(e.end - b) < 1e-3), None)
            return a, b, tk
        return NAN, NAN, None

    def snap_target(self, a, b, tk):
        if tk is not None or isnan(a) or not self.events:
            return a, b, tk
        tr = [i for i, e in enumerate(self.events) if e.kind == 'trans']
        if not tr:
            return a, b, tk
        ov = [min(b, self.events[i].end) - max(a, self.events[i].start) for i in tr]
        i = int(np.argmax(ov))
        m = ov[i]
        if m <= 0:
            return a, b, tk
        pl = self.events[tr[i]].end - self.events[tr[i]].start
        if m >= 0.8 * (b - a) and m >= 0.6 * pl:
            tk = tr[i]
            return self.events[tk].start, self.events[tk].end, tk
        return a, b, tk

    def push_undo(self):
        self.undo_stack.append(([replace(e) for e in self.events], self.next_uid))
        if len(self.undo_stack) > MAXUNDO:
            self.undo_stack.pop(0)

    def events_changed(self):
        self.dirty = True
        self.update_title()
        self.update_event_list()
        self.redraw_events()

    # ======================= labels ===========================================
    def disfl_name(self, abbr):
        return next((d['name'] for d in DISFL if d['abbr'] == abbr), 'unknown type')

    def err_name(self, abbr):
        return next((e['name'] for e in ERRS if e['abbr'] == abbr), 'unknown type')

    def kind_label(self, k):
        e = self.events[k]
        if e.kind == 'trans':
            return 'transcript "%s"' % e.transcript
        if e.kind == 'fluent':
            return 'fluent word "%s"' % e.transcript
        if e.kind == 'disfl':
            return '%s "%s"' % (e.cat, e.transcript)
        return '%s error' % e.cat

    def event_label(self, k):
        e = self.events[k]
        return '%s (%s - %s)' % (self.kind_label(k), fmt_time(e.start), fmt_time(e.end))

    def event_code(self, k):
        e = self.events[k]
        if e.kind == 'fluent':
            return 'F'
        if e.kind == 'disfl':
            d = next((d for d in DISFL if d['abbr'] == e.cat), None)
            return e.cat.upper() if d is None else d['key'].upper()
        if e.kind == 'error':
            return 'E-' + re.sub(r'^er-', '', e.cat).upper()
        return ''

    def fluent_notes(self, f):
        parts = []
        t = self.find_uid(self.events[f].tok)
        if t is not None and self.events[t].notes:
            parts.append(self.events[t].notes)
        if self.events[f].notes:
            parts.append(self.events[f].notes)
        return '; '.join(parts)

    # ======================= drawing ==========================================
    def _add(self, vb, item):
        vb.addItem(item, ignoreBounds=True)
        self.ev_items.append((vb, item))
        return item

    def _rect(self, vb, x0, x1, y0, y1, face=None, falpha=1.0, edge=None, lw=1.0):
        it = QtWidgets.QGraphicsRectItem(QRectF(x0, y0, x1 - x0, y1 - y0))
        if edge is None:
            it.setPen(QPen(Qt.NoPen))
        else:
            pen = QPen(qc(edge))
            pen.setCosmetic(True)
            pen.setWidthF(lw)
            it.setPen(pen)
        it.setBrush(QBrush(Qt.NoBrush) if face is None else QBrush(qc(face, falpha)))
        it.setZValue(1)
        return self._add(vb, it)

    def _text(self, vb, x, y, s, color=(0, 0, 0), size=10, bold=False, anchor=(0, 0.5), fill=None):
        if not s:
            return None
        ti = pg.TextItem(s, color=qc(color), anchor=anchor,
                         fill=None if fill is None else pg.mkBrush(qc(fill)))
        f = QFont()
        f.setPointSizeF(size)
        f.setBold(bold)
        ti.setFont(f)
        ti.setPos(x, y)
        ti.setZValue(3)
        return self._add(vb, ti)

    def fit_text(self, s, vb, ev, size):
        if not s:
            return s
        vw = max(self.viewE - self.viewS, 1e-9)
        w = (min(ev.end, self.viewE) - max(ev.start, self.viewS)) / vw * vb.width() - 6
        f = QFont()
        f.setPointSizeF(size)
        fm = QFontMetrics(f)
        if fm.horizontalAdvance(s) <= w:
            return s
        fw = s.split()[0] if s.split() else ''
        if fw and fm.horizontalAdvance(fw + '…') <= w:
            return fw + '…'
        if w < fm.horizontalAdvance('W…'):
            return ''
        return fm.elidedText(s, Qt.ElideRight, int(w))

    def clear_ev_items(self):
        for vb, it in self.ev_items:
            try:
                vb.removeItem(it)
            except Exception:
                pass
        self.ev_items = []

    def redraw_events(self):
        self.clear_ev_items()
        if not self.loaded:
            return
        self.compute_lanes()
        vS, vE = self.viewS, self.viewE
        vw = vE - vS
        fmax = self.max_freq
        selU = self.group_uid(self.cur) if (self.cur is not None and self.cur < len(self.events)) else 0
        for k, ev in enumerate(self.events):
            if ev.end < vS or ev.start > vE:
                continue
            col = self.event_color(k)
            is_sel = self.cur == k
            fa, lw = (0.40, 2.0) if is_sel else (0.22, 1.0)
            tx0 = max(ev.start, vS) + 0.003 * vw
            if ev.kind == 'trans':
                pc = self.piece_colors(k)
                if not pc:
                    self._rect(self.vb_trans, ev.start, ev.end, 0.04, 0.96, (0.96,) * 3, 1, (0.35,) * 3, lw)
                else:
                    hb = 0.92 / len(pc)
                    for q, c in enumerate(pc):
                        y0 = 0.96 - (q + 1) * hb
                        self._rect(self.vb_trans, ev.start, ev.end, y0, y0 + hb, mix(c, 0.55, 0.45))
                    self._rect(self.vb_trans, ev.start, ev.end, 0.04, 0.96, None, 1, (0.35,) * 3, lw)
                att = self.attached_to(k)
                is_flu = any(self.events[j].kind == 'fluent' for j in att)
                in_sel = selU > 0 and any(self.events[j].uid == selU or self.events[j].link == selU
                                          for j in att)
                tc = (0.85, 0, 0) if in_sel else ((0.10, 0.30, 0.85) if is_flu else (0, 0, 0))
                self._text(self.vb_trans, tx0, 0.5, self.fit_text(ev.transcript, self.vb_trans, ev, 10),
                           tc, 10, is_flu or in_sel)
            elif ev.kind == 'fluent':
                pass
            else:
                if ev.kind == 'error':
                    self._rect(self.vb_spec, ev.start, ev.end, fmax * 0.88, fmax, col, 0.55, col, lw)
                    self._rect(self.vb_wave, ev.start, ev.end, -1, -0.76, col, 0.55, col, lw)
                    faD = 0.35
                else:
                    self._rect(self.vb_spec, ev.start, ev.end, 0, fmax, col, fa,
                               tuple(c * 0.8 for c in col), lw)
                    self._rect(self.vb_wave, ev.start, ev.end, -1, 1, col, fa,
                               tuple(c * 0.8 for c in col), lw)
                    faD = max(fa, 0.35)
                b0, b1 = self.lane_band(k)
                self._rect(self.vb_de, ev.start, ev.end, b0, b1, col, faD, tuple(c * 0.8 for c in col), lw)
                fsz = 7 if (b1 - b0) < 0.11 else 8
                self._text(self.vb_de, tx0, (b0 + b1) / 2, self.fit_text(ev.cat, self.vb_de, ev, fsz),
                           (0, 0, 0), fsz, True)
            if ev.notes:
                ty = 0.78 if ev.kind == 'trans' else (0.22 if ev.kind == 'error' else 0.5)
                self._rect(self.vb_notes, ev.start, ev.end, ty - 0.2, ty + 0.2, None, 1, col, lw)
                self._text(self.vb_notes, tx0, ty, self.fit_text(ev.notes, self.vb_notes, ev, 8),
                           (0, 0, 0), 8)
        for k in self.focus_set():
            ev = self.events[k]
            if ev.end < vS or ev.start > vE:
                continue
            if ev.kind == 'trans':
                self._rect(self.vb_trans, ev.start, ev.end, 0, 1, None, 1, (0, 0, 0), 3)
            elif ev.kind != 'fluent':
                b0, b1 = self.lane_band(k)
                self._rect(self.vb_de, ev.start, ev.end, b0, b1, None, 1, (0, 0, 0), 2.5)
        if self.group_u > 0:
            f = self.find_uid(self.group_u)
            if f is not None and self.events[f].end > vS and self.events[f].start < vE:
                e = self.events[f]
                self._rect(self.vb_spec, e.start, e.end, 0, fmax, FLUENT_COLOR, 0.25, (0, 0, 0), 2.5)
                self._rect(self.vb_wave, e.start, e.end, -1, 1, FLUENT_COLOR, 0.25, (0, 0, 0), 2.5)
                xm = (max(e.start, vS) + min(e.end, vE)) / 2
                self._text(self.vb_spec, xm, fmax * 0.5, 'fluent event selected', (0, 0, 0), 10, True,
                           (0.5, 0.5), (1, 1, 1))
                self._text(self.vb_wave, xm, 0, 'highlight a part + Enter  |  Esc = done', (0, 0, 0), 8,
                           False, (0.5, 0.5), (1, 1, 1))
        self.draw_overview_ticks()
        self.position_box()
        self.update_annot_display()

    def update_selection_graphics(self):
        self.update_info()
        show = self.has_sel() and self.loaded
        for vb, r, l1, l2 in self.sel_items:
            if not show:
                r.hide()
                l1.hide()
                l2.hide()
                continue
            y0, y1 = vb.yr
            r.setRect(QRectF(self.sel_s, y0, self.sel_e - self.sel_s, y1 - y0))
            l1.setPos(self.sel_s)
            l2.setPos(self.sel_e)
            r.show()
            l1.show()
            l2.show()
        self.sync_play_btn()

    def update_cursor_graphics(self):
        for c in self.cur_lines:
            c.setPos(self.cursor if not isnan(self.cursor) else -1)
        self.ov_cur.setPos(self.cursor if not isnan(self.cursor) else -1)
        self.update_info()

    def sync_play_btn(self):
        self.btn_play.setText('Pause (Space)' if self.playing else 'Play (Space)')

    def update_info(self):
        if not self.has_sel():
            s = 'Cursor: %s\nNo selection - drag on the timeline.' % fmt_time(self.cursor)
        else:
            s = 'Selection: %s - %s\nLength: %.3f s' % (fmt_time(self.sel_s), fmt_time(self.sel_e),
                                                       self.sel_e - self.sel_s)
            if self.cur is not None and self.cur < len(self.events):
                s += '   |   ' + self.kind_label(self.cur)
        self.info.setText(s)

    def update_annot_display(self):
        if self.pending_new is not None:
            a, b = self.pending_new
            L = ['NEW TRANSCRIPT', '%s - %s' % (fmt_time(a), fmt_time(b)), '',
                 'Type what was said, e.g.  t-t-table', 'It is split into pieces you can',
                 'right-click and mark.', '', 'Enter = add,  Esc = cancel']
        elif self.cur is None or self.cur >= len(self.events):
            L = ['Nothing selected.', '', 'Highlight speech -> double-click -> type it.',
                 'Highlight a word -> Enter = fluent event.', 'Right-click (two-finger click) a piece',
                 'to mark it fluent / disfluent / error.']
        else:
            k = self.cur
            ev = self.events[k]
            L = []
            if ev.kind == 'trans':
                if any(self.events[j].kind == 'fluent' for j in self.attached_to(k)):
                    L.append('FLUENT EVENT  "%s"' % ev.transcript)
                else:
                    L.append('PIECE  "%s"' % ev.transcript)
            elif ev.kind == 'fluent':
                L.append('FLUENT WORD  "%s"' % ev.transcript)
            elif ev.kind == 'disfl':
                L.append('DISFLUENCY:  %s  (%s)' % (ev.cat, self.disfl_name(ev.cat)))
            else:
                L.append('ERROR:  %s  (%s)' % (ev.cat, self.err_name(ev.cat)))
            L.append('%s - %s   (%.3f s)' % (fmt_time(ev.start), fmt_time(ev.end), ev.end - ev.start))
            a, b, _ = self.current_target()
            L.append('Marked as: ' + self.piece_types(a, b))
            L.append('')
            if ev.kind == 'disfl':
                L.append('What was said: "%s"' % ev.transcript)
                f = self.find_uid(ev.link)
                L.append('Resolves to: (no fluent word yet)' if f is None else
                         'Resolves to: "%s"' % self.events[f].transcript)
                L.append('')
            L.append('Notes:  (N to edit)')
            L.append(ev.notes or '(none)')
            u = self.group_uid(k)
            if u > 0:
                L = self.group_lines(u) + L
        if self.group_u > 0 and self.pending_new is None:
            f = self.find_uid(self.group_u)
            if f is not None:
                L = ['>> FLUENT EVENT SELECTED: "%s"' % self.events[f].transcript,
                     '   Highlight a disfluent part + Enter', '   W = whole episode,  Esc = done',
                     '   New speech: highlight + double-click', ''] + L
        self.annot.setPlainText('\n'.join(L))

    def group_lines(self, u):
        f = self.find_uid(u)
        if f is None:
            return []
        m = sorted([f] + self.linked_all(u), key=lambda i: self.events[i].start)
        L = ['"%s" BREAKDOWN' % self.events[f].transcript]
        for q in m:
            e = self.events[q]
            if e.kind == 'fluent':
                lab, tx = 'fluent', e.transcript
            elif e.kind == 'disfl':
                lab, tx = e.cat, e.transcript
            else:
                lab, tx = e.cat, '(' + self.err_name(e.cat) + ')'
            mark = '> ' if q == self.cur else '  '
            L.append('%s%s  %-6s %s' % (mark, fmt_time(e.start), lab, tx))
        L.append('')
        return L

    def draw_overview_ticks(self):
        for it in self.ov_items:
            self.ov_vb.removeItem(it)
        self.ov_items = []
        if not self.loaded:
            return
        self.ov_rect.setRect(QRectF(self.viewS, 0, self.viewE - self.viewS, 1))
        for d in DISFL:
            ev = [e for e in self.events if e.kind == 'disfl' and e.cat == d['abbr']]
            if ev:
                x = np.array([[e.start, e.end] for e in ev]).ravel()
                c = pg.PlotCurveItem(x, np.full(len(x), 0.12), pen=pg.mkPen(qc(d['color']), width=4),
                                     connect='pairs')
                self.ov_vb.addItem(c, ignoreBounds=True)
                self.ov_items.append(c)
        ev = [e for e in self.events if e.kind == 'error']
        if ev:
            x = np.array([[e.start, e.end] for e in ev]).ravel()
            c = pg.PlotCurveItem(x, np.full(len(x), 0.88), pen=pg.mkPen(qc(ERR_COLOR), width=4),
                                 connect='pairs')
            self.ov_vb.addItem(c, ignoreBounds=True)
            self.ov_items.append(c)

    def build_overview(self):
        n = len(self.audio)
        nb = min(2000, n)
        L = (n // nb) * nb
        env = np.abs(self.audio[:L]).reshape(nb, -1).max(axis=1) / self.amax
        tc = (np.arange(nb) + 0.5) * (L / nb) / self.fs
        x = np.repeat(tc, 2)
        y = np.column_stack([0.5 + 0.45 * env, 0.5 - 0.45 * env]).ravel()
        self.ov_env.setData(x, y, connect='pairs')
        self.ov_vb.setXRange(0, max(self.dur, 1e-6), padding=0)
        self.ov_vb.setYRange(0, 1, padding=0)

    def draw_spectrogram(self):
        fs, a = self.fs, self.audio
        i0 = max(0, int(math.floor(self.viewS * fs)))
        i1 = min(len(a), int(math.ceil(self.viewE * fs)))
        win = max(64, int(round(0.006 * fs)))
        if win % 2 == 0:
            win += 1
        if i1 - i0 < max(64, win):
            self.spec_img.clear()
            return
        seg = a[i0:i1].astype(np.float64)
        ncol = 700
        hop = max(1, (len(seg) - win) // ncol)
        nfft = 2 ** int(math.ceil(math.log2(max(win, 512))))
        w = 0.5 - 0.5 * np.cos(2 * np.pi * np.arange(win) / (win - 1))
        st = np.arange(0, len(seg) - win + 1, hop)
        frames = np.lib.stride_tricks.sliding_window_view(seg, win)[st]
        X = np.fft.rfft(frames * w, nfft, axis=1)
        P = 20 * np.log10(np.abs(X) + 1e-12)
        F = np.arange(P.shape[1]) * fs / nfft
        fm = F <= self.max_freq
        if not fm.any():
            fm[:] = True
        P = P[:, fm].T.astype(np.float32)          # rows = frequency, cols = time
        self.spec_max = float(P.max())
        df = fs / nfft
        x0 = (i0 + st[0] + win / 2 - hop / 2) / fs
        x1 = (i0 + st[-1] + win / 2 + hop / 2) / fs
        self.spec_img.setImage(P, autoLevels=False)
        self.spec_img.setRect(QRectF(x0, 0, max(x1 - x0, 1e-6), F[fm][-1] + df))
        self.apply_contrast()

    def apply_contrast(self):
        if not isnan(self.spec_max):
            self.spec_img.setLevels([self.spec_max - (100 - self.contrast), self.spec_max])

    def set_contrast(self, v):
        self.contrast = int(min(80, max(0, round(v))))
        if self.sld.value() != self.contrast:
            self.sld.blockSignals(True)
            self.sld.setValue(self.contrast)
            self.sld.blockSignals(False)
        self.apply_contrast()
        self.settings.setValue('specContrast', self.contrast)
        self.set_status('Spectrogram contrast %d / 80  ([ less, ] more).' % self.contrast)

    def draw_waveform(self):
        fs = self.fs
        i0 = max(0, int(math.floor(self.viewS * fs)))
        i1 = min(len(self.audio), int(math.ceil(self.viewE * fs)))
        if i1 <= i0 + 1:
            self.wave_curve.setData([], [])
            return
        seg = self.audio[i0:i1]
        if len(seg) <= 4000:
            tt = (i0 + np.arange(len(seg))) / fs
            self.wave_curve.setData(tt, seg / self.amax, connect='all')
        else:
            nb = 1500
            L = (len(seg) // nb) * nb
            M = seg[:L].reshape(nb, -1)
            tc = (i0 + np.arange(nb) * (L / nb)) / fs
            y = np.column_stack([M.max(axis=1), M.min(axis=1)]).ravel() / self.amax
            self.wave_curve.setData(np.repeat(tc, 2), y, connect='pairs')

    def refresh_view(self):
        if not self.loaded:
            return
        self.viewS = max(0.0, self.viewS)
        self.viewE = min(self.dur, self.viewE)
        if self.viewE - self.viewS < MINVIEW:
            self.viewE = min(self.dur, self.viewS + MINVIEW)
        self.vb_spec.setXRange(self.viewS, self.viewE, padding=0)
        for fn in (self.draw_spectrogram, self.draw_waveform, self.redraw_events):
            try:
                fn()
            except Exception as err:
                self.set_status('Draw error: %s' % err, 'warn')
        self.update_selection_graphics()
        self.update_cursor_graphics()
        self.view_lbl.setText('View: %s - %s  (%.2f s wide)' % (fmt_time(self.viewS), fmt_time(self.viewE),
                                                               self.viewE - self.viewS))
        self.sync_time_fields(self.viewS, self.viewE)

    # ======================= view / zoom ======================================
    def clamp_t(self, t):
        return max(0.0, min(self.dur, t))

    def set_view(self, a, b):
        if not self.loaded:
            return
        w = b - a
        if w < MINVIEW:
            c = (a + b) / 2
            a, b, w = c - MINVIEW / 2, c + MINVIEW / 2, MINVIEW
        if w > self.dur:
            a, b = 0, self.dur
        if a < 0:
            b, a = b - a, 0
        if b > self.dur:
            a, b = a - (b - self.dur), self.dur
        self.viewS, self.viewE = max(0.0, a), min(self.dur, b)
        self.refresh_view()

    def zoom_about(self, center, factor):
        if not self.loaded:
            return
        if isnan(center):
            center = (self.viewS + self.viewE) / 2
        w = max(MINVIEW, min(self.dur, (self.viewE - self.viewS) * factor))
        frac = (center - self.viewS) / max(self.viewE - self.viewS, 1e-12)
        frac = max(0, min(1, frac))
        a = center - frac * w
        self.set_view(a, a + w)

    def full_view(self):
        self.set_view(0, self.dur)

    def zoom_to_selection(self):
        if not self.has_sel():
            QtWidgets.QMessageBox.information(self, 'Fit selection',
                                              'There is no selection to zoom to.\n\nClick and drag across the '
                                              'spectrogram or waveform first (or click an existing event).')
            return
        pad = 0.1 * (self.sel_e - self.sel_s)
        self.set_view(self.sel_s - pad, self.sel_e + pad)

    def scroll_by(self, frac):
        sh = frac * (self.viewE - self.viewS)
        self.set_view(self.viewS + sh, self.viewE + sh)

    def cursor_center(self):
        return (self.viewS + self.viewE) / 2 if isnan(self.cursor) else self.cursor

    def zoom_center(self):
        return (self.sel_s + self.sel_e) / 2 if self.has_sel() else self.cursor_center()

    def on_wheel(self, delta, mods, t):
        if self.editing:
            return
        c = -delta / 120.0
        if mods & (Qt.ControlModifier | Qt.MetaModifier):
            self.zoom_about(t, 1.2 ** c)
        else:
            self.scroll_by(0.15 * c)

    def sync_time_fields(self, a, b, force=False):
        if force or not self.edt_from.hasFocus():
            self.edt_from.setText(fmt_time(a))
        if force or not self.edt_to.hasFocus():
            self.edt_to.setText(fmt_time(b))

    def apply_time_fields(self, changed=''):
        if not self.loaded:
            return
        a, b = parse_time(self.edt_from.text()), parse_time(self.edt_to.text())
        self.glw.setFocus()
        if isnan(a) or isnan(b):
            QtWidgets.QMessageBox.warning(self, 'Invalid time',
                                          'That time could not be read.\n\nType seconds (e.g. 14 or 75.5), '
                                          'minutes:seconds (e.g. 1:05), or 2m30.')
            self.refresh_view()
            return
        a, b = min(max(a, 0), self.dur), min(max(b, 0), self.dur)
        if b <= a:
            w = max(self.viewE - self.viewS, MINVIEW)
            if changed == 'to':
                a = max(0, b - w)
            else:
                b = min(self.dur, a + w)
        if b <= a:
            QtWidgets.QMessageBox.warning(self, 'Invalid time frame',
                                          'The start time is at or past the end of the recording (%s).'
                                          % fmt_time(self.dur))
            self.refresh_view()
            return
        self.set_view(a, b)
        self.select_whole_view()

    def select_whole_view(self):
        self.sel_s, self.sel_e = self.viewS, self.viewE
        self.cur = None
        self.cursor = self.sel_s
        self.redraw_events()
        self.update_selection_graphics()
        self.update_cursor_graphics()
        self.show_frame_at(self.cursor)
        self.set_status(self.sel_hint(), 'ok')

    def sel_hint(self):
        if self.group_u > 0:
            return ('Highlighted %s - %s. Enter = disfluency / error type(s) | double-click = new transcript '
                    '| Esc = done with this word.' % (fmt_time(self.sel_s), fmt_time(self.sel_e)))
        return ('Highlighted %s - %s (%.3f s). Double-click = type transcript | Enter = fluent word (if on the '
                'transcript) | Space = play | right-click = menu.'
                % (fmt_time(self.sel_s), fmt_time(self.sel_e), self.sel_e - self.sel_s))

    # ---- overview ----
    def ov_hit(self, t):
        tol = 6 * max(self.dur, 1e-9) / max(self.ov_vb.width(), 1)
        if abs(t - self.viewS) <= tol:
            return 'left'
        if abs(t - self.viewE) <= tol:
            return 'right'
        if self.viewS < t < self.viewE:
            return 'pan'
        return ''

    def ov_click(self, t):
        self.commit_edit()
        if not self.ov_hit(t):
            w = self.viewE - self.viewS
            self.set_view(t - w / 2, t + w / 2)

    def ov_start(self, t):
        self.commit_edit()
        self.ov_mode = self.ov_hit(t)
        if not self.ov_mode:
            w = self.viewE - self.viewS
            self.set_view(t - w / 2, t + w / 2)
            self.ov_mode = 'pan'
        self.ov_grab = t - self.viewS

    def ov_drag(self, t):
        if self.ov_mode == 'left':
            self.set_view(min(t, self.viewE - MINVIEW), self.viewE)
        elif self.ov_mode == 'right':
            self.set_view(self.viewS, max(t, self.viewS + MINVIEW))
        else:
            w = self.viewE - self.viewS
            a = t - self.ov_grab
            self.set_view(a, a + w)

    def ov_end(self):
        self.ov_mode = ''
        self.set_status('Viewing %s - %s.' % (fmt_time(self.viewS), fmt_time(self.viewE)))

    # ======================= mouse on the timelines ===========================
    def hit_edge(self, vb, t):
        if not self.has_sel():
            return ''
        tol = EDGE_PX * (self.viewE - self.viewS) / max(vb.width(), 1)
        dL, dR = abs(t - self.sel_s), abs(t - self.sel_e)
        if min(dL, dR) > tol:
            return ''
        return 'left' if dL <= dR else 'right'

    def on_hover(self, vb, t, y):
        if self.editing:
            return
        if self.hit_edge(vb, t):
            cur = Qt.SizeHorCursor
        elif self.event_at_time(t, vb.panel, y) is not None:
            cur = Qt.SizeAllCursor
        else:
            cur = Qt.ArrowCursor
        if self.glw.viewport().cursor().shape() != cur:
            self.glw.viewport().setCursor(cur)

    def on_click(self, panel, t, y):
        self.commit_edit()
        idx = self.event_at_time(t, panel, y)
        if idx is not None:
            self.cur = idx
            self.sel_s, self.sel_e = self.events[idx].start, self.events[idx].end
            self.cursor = self.sel_s
            self.sync_list_selection()
            self.set_status('Selected %s. Right-click (two-finger click) for options, double-click to edit its '
                            'text, Space to play.' % self.event_label(idx))
        elif panel in ('', 'transcript') and self.has_sel() and self.cur is None and self.sel_s < t < self.sel_e:
            self.cursor = t
            self.set_status('Play point %s - Space plays from here. Double-click the highlight to type its '
                            'transcript.' % fmt_time(t))
        elif panel == '':
            self.sel_s, self.sel_e = NAN, NAN
            self.cur = None
            self.cursor = t
        self.redraw_events()
        self.update_selection_graphics()
        self.update_cursor_graphics()
        self.show_frame_at(self.cursor)
        self.glw.setFocus()

    def on_double(self, panel, t, y):
        self.commit_edit()
        in_hl = self.has_sel() and self.sel_s <= t <= self.sel_e
        if in_hl:
            if panel == 'notes':
                if self.cur is not None and self.cur < len(self.events):
                    k = self.cur
                else:
                    k = self.piece_under(self.sel_s, self.sel_e, t)
                    if k is None:
                        k = self.event_at_time(t, 'notes', y)
                if k is None:
                    self.set_status('Nothing here to attach notes to yet - double-click the highlight on the '
                                    'spectrogram or Transcript strip to type its transcript first.', 'warn')
                    return
                self.start_editing(k, 'notes')
                return
            if self.group_u > 0:
                self.exit_group()
            if not self.overlap_idx(self.sel_s, self.sel_e, ['trans']):
                self.begin_new_text()
            else:
                k = self.piece_under(self.sel_s, self.sel_e, t)
                if k is not None:
                    self.start_editing(k, 'text')
            return
        idx = self.event_at_time(t, panel, y)
        if idx is None or not panel:
            return
        self.start_editing(idx, 'notes' if panel == 'notes' else 'text')

    def on_right(self, panel, t, y, screen_pos):
        self.commit_edit()
        self.right_click_pick(t, panel, y)
        self.show_context_menu(screen_pos.toPoint() if hasattr(screen_pos, 'toPoint') else screen_pos)

    def right_click_pick(self, t, panel, y):
        if self.has_sel() and self.cur is None and self.sel_s <= t <= self.sel_e:
            return
        idx = self.event_at_time(t, panel, y)
        if idx is not None and idx != self.cur:
            self.cur = idx
            self.sel_s, self.sel_e = self.events[idx].start, self.events[idx].end
            self.cursor = self.sel_s
            self.sync_list_selection()
            self.redraw_events()
            self.update_selection_graphics()
            self.update_cursor_graphics()

    def drag_start(self, vb, panel, t, y, mods):
        self.commit_edit()
        shift = bool(mods & Qt.ShiftModifier)
        d = SimpleNamespace(vb=vb, panel=panel, t0=t, mode='', move_idx=None, orig=(0, 0), kids=[],
                            kids_orig=[], pushed=False, nbr=None)
        if not shift:
            d.mode = self.hit_edge(vb, t)
        if d.mode:
            if self.cur is not None:
                self.push_undo()
                if self.events[self.cur].kind == 'trans':
                    d.nbr = self.touching_piece(self.cur, d.mode)
        else:
            hit = self.event_at_time(t, panel, y)
            if panel in ('', 'transcript') and hit is not None:
                d.mode = 'select'
            elif not shift and hit is not None:
                d.mode = 'move'
                d.move_idx = hit
                d.orig = (self.events[hit].start, self.events[hit].end)
                if self.events[hit].kind == 'trans':
                    d.kids = self.attached_to(hit)
                    d.kids_orig = [(self.events[j].start, self.events[j].end) for j in d.kids]
            else:
                d.mode = 'select'
        self.drag = d

    def drag_move(self, t):
        d = self.drag
        if d is None:
            return
        E = self.events
        if d.mode == 'left':
            lo = 0.0 if d.nbr is None else E[d.nbr].start + MINSEL
            self.sel_s = max(lo, min(t, self.sel_e - MINSEL))
            if self.cur is not None:
                E[self.cur].start = self.sel_s
                self.clear_global(self.cur)
                if E[self.cur].kind == 'trans':
                    self.sync_tok(self.cur)
                else:
                    E[self.cur].tok = 0
                if d.nbr is not None:
                    E[d.nbr].end = self.sel_s
                    self.clear_global(d.nbr)
                    self.sync_tok(d.nbr)
                self.redraw_events()
        elif d.mode == 'right':
            hi = self.dur if d.nbr is None else E[d.nbr].end - MINSEL
            self.sel_e = min(hi, max(t, self.sel_s + MINSEL))
            if self.cur is not None:
                E[self.cur].end = self.sel_e
                self.clear_global(self.cur)
                if E[self.cur].kind == 'trans':
                    self.sync_tok(self.cur)
                else:
                    E[self.cur].tok = 0
                if d.nbr is not None:
                    E[d.nbr].start = self.sel_e
                    self.clear_global(d.nbr)
                    self.sync_tok(d.nbr)
                self.redraw_events()
        elif d.mode == 'move':
            if not d.pushed:
                self.push_undo()
                d.pushed = True
                self.cur = d.move_idx
            a = d.orig[0] + (t - d.t0)
            b = d.orig[1] + (t - d.t0)
            if a < 0:
                b, a = b - a, 0
            if b > self.dur:
                a, b = a - (b - self.dur), self.dur
            dd = a - d.orig[0]
            self.sel_s, self.sel_e = a, b
            k = d.move_idx
            E[k].start, E[k].end = a, b
            self.clear_global(k)
            if E[k].kind != 'trans':
                E[k].tok = 0
            for j, (s0, e0) in zip(d.kids, d.kids_orig):
                E[j].start, E[j].end = s0 + dd, e0 + dd
                self.clear_global(j)
            self.redraw_events()
        else:
            self.sel_s, self.sel_e = min(d.t0, t), max(d.t0, t)
            self.cur = None
        self.update_selection_graphics()
        if not isnan(self.sel_s):
            self.sync_time_fields(self.sel_s, self.sel_e)

    def drag_finish(self):
        d = self.drag
        self.drag = None
        if d is None:
            return
        if d.mode in ('left', 'right'):
            self.cursor = self.sel_s
            if self.cur is not None:
                if not self.partition_ok(self.cur):
                    self.revert_last('Not allowed: overlap - resize undone.')
                else:
                    self.relink_all()
                    self.dirty = True
                    self.update_title()
                    self.update_event_list()
                    self.redraw_events()
                    self.set_status('Resized %s.' % self.event_label(self.cur), 'ok')
            else:
                self.set_status(self.sel_hint())
        elif d.mode == 'move':
            self.cursor = self.sel_s
            if d.move_idx is not None:
                if not self.partition_ok(d.move_idx):
                    self.revert_last('Not allowed: overlap - move undone.')
                else:
                    self.relink_all()
                    self.dirty = True
                    self.update_title()
                    self.update_event_list()
                    self.redraw_events()
                    self.set_status('Moved %s (Ctrl+Z to undo).' % self.event_label(d.move_idx), 'ok')
        else:
            self.cursor = self.sel_s
            self.redraw_events()
            self.set_status(self.sel_hint())
        self.update_selection_graphics()
        self.update_cursor_graphics()
        self.show_frame_at(self.cursor)
        self.glw.setFocus()

    def revert_last(self, msg):
        if not self.undo_stack:
            return
        self.events, self.next_uid = self.undo_stack.pop()
        if self.cur is not None and self.cur < len(self.events):
            self.sel_s, self.sel_e = self.events[self.cur].start, self.events[self.cur].end
        else:
            self.cur = None
        self.events_changed()
        self.update_selection_graphics()
        self.set_status(msg, 'warn')
        QtWidgets.QApplication.beep()

    def step_cursor(self, direction, big):
        if self.playing:
            return
        step = 1.0 if big else 1.0 / max(self.vsrc.fps, 1)
        t0 = self.viewS if isnan(self.cursor) else self.cursor
        self.cursor = self.clamp_t(t0 + direction * step)
        if self.cursor < self.viewS or self.cursor > self.viewE:
            w = self.viewE - self.viewS
            self.set_view(self.cursor - w / 2, self.cursor + w / 2)
        self.update_cursor_graphics()
        self.show_frame_at(self.cursor)
        self.set_status('Cursor at %s' % fmt_time(self.cursor))

    def clear_selection(self):
        self.sel_s, self.sel_e = NAN, NAN
        self.cur = None
        self.redraw_events()
        self.update_selection_graphics()
        self.set_status('Selection cleared.')

    # ======================= context menu =====================================
    def show_context_menu(self, pos):
        a, b, tk = self.current_target()
        m = QtWidgets.QMenu(self)
        if isnan(a):
            act = m.addAction('Highlight some speech or click a transcript piece first')
            act.setEnabled(False)
            m.exec(pos)
            return

        def add(label, checked, fn, enabled=True):
            act = m.addAction(label)
            if checked is not None:
                act.setCheckable(True)
                act.setChecked(checked)
            act.setEnabled(enabled)
            act.triggered.connect(lambda: self.cb(fn))
            return act

        add('Fluent', self.has_annot(a, b, 'fluent', '') is not None, lambda: self.menu_toggle('fluent', ''))
        for d in DISFL:
            add('%s  (%s)' % (d['name'], d['abbr']), self.has_annot(a, b, 'disfl', d['abbr']) is not None,
                lambda ab=d['abbr']: self.menu_toggle('disfl', ab))
        em = m.addMenu('Error')
        for e in ERRS:
            act = em.addAction('%s  -  %s' % (e['abbr'], e['name']))
            act.setCheckable(True)
            act.setChecked(self.has_annot(a, b, 'error', e['abbr']) is not None)
            act.triggered.connect(lambda _=False, ab=e['abbr']: self.cb(lambda: self.menu_toggle('error', ab)))
        add('Several...', None, self.open_event_picker)
        m.addSeparator()
        can_split = tk is not None and len(split_tokens(self.events[tk].transcript)) > 1
        add('Split into words', None, self.menu_split, can_split)
        lab = 'Edit text  ("%s")' % self.events[tk].transcript if tk is not None else 'Edit text'
        add(lab, None, self.edit_current_event, self.cur is not None)
        add('Play', None, self.start_play_fresh)
        add('Delete', None, self.delete_current_event, self.cur is not None)
        m.exec(pos)

    def menu_toggle(self, kind, cat):
        a, b, tk = self.current_target()
        if isnan(a):
            return
        j = self.has_annot(a, b, kind, cat)
        if j is None:
            self.add_annotation(kind, cat)
            return
        keep = self.events[tk].uid if tk is not None else 0
        self.push_undo()
        del self.events[j]
        self.relink_all()
        self.cur = self.find_uid(keep) if keep else None
        self.events_changed()
        self.update_selection_graphics()
        self.set_status('Removed it. This piece now has: %s.' % self.piece_types(a, b), 'ok')

    def add_annotation(self, kind, cat):
        a, b, tk = self.current_target()
        if isnan(a):
            self.set_status('Highlight some speech or click a transcript piece first.', 'warn')
            return None
        if self.has_annot(a, b, kind, cat) is not None:
            self.set_status('This piece already has that. It has: %s.' % self.piece_types(a, b))
            return None
        if kind == 'fluent':
            ov = self.overlap_idx(a, b, ['fluent'])
            if ov:
                self.set_status('That overlaps the fluent word "%s".' % self.events[ov[0]].transcript, 'warn')
                QtWidgets.QApplication.beep()
                return None
        self.push_undo()
        k = self.add_raw(a, b, tk, kind, cat)
        self.relink_all()
        self.cur = tk if tk is not None else k
        self.sel_s, self.sel_e = a, b
        self.events_changed()
        self.update_selection_graphics()
        self.set_status('This piece has: %s. Right-click again to add more (Ctrl+Z to undo).'
                        % self.piece_types(a, b), 'ok')
        if tk is None and DISFL_HAS_TRANSCRIPT and kind != 'error' and cat != 'bl':
            self.start_editing(k, 'text')
        return k

    def menu_split(self):
        _, _, tk = self.current_target()
        if tk is None:
            return
        ev = self.events[tk]
        if len(split_tokens(ev.transcript)) < 2:
            return
        self.push_undo()
        for j in self.attached_to(tk):
            self.events[j].tok = 0
        del self.events[tk]
        k1 = self.add_pieces(ev.start, ev.end, ev.transcript)
        self.cur = k1
        self.sel_s, self.sel_e = self.events[k1].start, self.events[k1].end
        self.events_changed()
        self.update_selection_graphics()
        self.set_status('Split into pieces - drag the edges between them to fit the audio.', 'ok')

    # ======================= Enter workflow ===================================
    def on_enter(self):
        if self.group_u > 0 and self.find_uid(self.group_u) is None:
            self.group_u = 0
        a, b, tk = self.current_target()
        if isnan(a):
            if self.group_u > 0:
                self.set_status('Highlight a disfluent part (e.g. t-) and press Enter - or Esc when done.')
            else:
                self.set_status('Drag to highlight first: a word in the transcript, or new speech to transcribe.',
                                'warn')
            return
        if self.group_u > 0:
            f = self.find_uid(self.group_u)
            if abs(a - self.events[f].start) < 1e-3 and abs(b - self.events[f].end) < 1e-3:
                self.set_status('That is the fluent word itself. Highlight a disfluent part (e.g. t-) and press '
                                'Enter, W for the whole episode, or Esc when done.')
                return
            self.open_event_picker()
            return
        if not self.overlap_idx(a, b, ['trans']):
            if not self.has_sel():
                self.sel_s, self.sel_e = a, b
            self.begin_new_text()
            return
        if self.cur is not None and self.cur < len(self.events) and \
                self.events[self.cur].kind in ('disfl', 'error'):
            self.open_event_picker()
            return
        self.fluent_from_target()

    def begin_new_text(self):
        if not self.has_sel():
            QtWidgets.QMessageBox.information(self, 'Select first',
                                              'Drag across the spectrogram / waveform to select first.')
            return
        a, b = self.sel_s, self.sel_e
        ov = self.overlap_idx(a, b, ['trans'])
        if ov:
            self.set_status('That overlaps the transcript "%s" - click it on the Transcript strip to edit it, '
                            'or trim the selection.' % self.events[ov[0]].transcript, 'warn')
            QtWidgets.QApplication.beep()
            return
        if self.editing:
            self.commit_edit()
        self.group_u = 0
        self.cur = None
        self.pending_new = (a, b)
        self.box_target = -1
        self.box_kind = 'transcript'
        self.box.setText('')
        self.editing = True
        self.redraw_events()
        self.box.setFocus()
        self.set_status('TRANSCRIPT: type what was said, Enter to add it, Esc to cancel.')

    def on_part_key(self, code):
        if code == 'flu':
            self.fluent_from_target()
            return
        a, b, tk = self.current_target()
        if isnan(a):
            self.set_status('Highlight the part first.', 'warn')
            return
        a, b, tk = self.snap_target(a, b, tk)
        self.sel_s, self.sel_e = a, b
        if tk is not None:
            self.cur = tk
        elif self.cur is not None and self.events[self.cur].kind not in ('disfl', 'error'):
            self.cur = None
        self.add_annotation('disfl', code)

    def start_add_error(self):
        a, _, _ = self.current_target()
        if isnan(a):
            self.set_status('Highlight some speech or click a transcript piece first.', 'warn')
            return
        q = self.error_picker('Add error')
        if q is None:
            self.set_status('Cancelled - no error added.')
            return
        self.add_annotation('error', ERRS[q]['abbr'])

    def error_picker(self, title):
        d = ErrorPicker(self, title)
        if d.exec() == QtWidgets.QDialog.Accepted:
            return d.choice
        return None

    def fluent_from_target(self):
        a, b, tk = self.current_target()
        if isnan(a):
            self.set_status('Highlight a word in the transcript first.', 'warn')
            return
        a, b, tk = self.snap_target(a, b, tk)
        j = self.has_annot(a, b, 'fluent', '')
        if j is not None:
            self.enter_group(self.events[j].uid)
            return
        ov = self.overlap_idx(a, b, ['fluent'])
        if ov:
            self.enter_group(self.events[ov[0]].uid)
            self.sel_s, self.sel_e = a, b
            self.cur = None
            self.update_selection_graphics()
            self.redraw_events()
            self.open_event_picker()
            return
        self.push_undo()
        k = self.add_raw(a, b, tk, 'fluent', '')
        if tk is None:
            self.events[k].transcript = self.words_under(a, b)
        self.relink_all()
        self.events_changed()
        self.enter_group(self.events[k].uid)

    def enter_group(self, u):
        self.group_u = u
        f = self.find_uid(u)
        tk = self.find_uid(self.events[f].tok)
        self.cur = tk if tk is not None else f
        self.sel_s, self.sel_e = NAN, NAN
        self.update_selection_graphics()
        self.redraw_events()
        self.sync_list_selection()
        self.set_status('FLUENT EVENT SELECTED: "%s". Now highlight each disfluent part (e.g. t-) and press Enter, '
                        'or W for the whole episode. Esc when done.' % self.events[f].transcript, 'ok')

    def exit_group(self):
        if self.group_u > 0:
            f = self.find_uid(self.group_u)
            self.group_u = 0
            if f is not None:
                self.set_status('Finished "%s".' % self.events[f].transcript, 'ok')
        self.redraw_events()

    def whole_episode(self):
        if self.group_u <= 0:
            self.set_status('W works after a fluent word is selected (highlight it + Enter).')
            return
        a, b = self.group_span(self.group_u)
        self.sel_s, self.sel_e = a, b
        self.cur = None
        self.cursor = a
        self.update_selection_graphics()
        self.redraw_events()
        self.open_event_picker()

    def open_event_picker(self):
        a, b, tk = self.current_target()
        if isnan(a):
            self.set_status('Highlight the part first.', 'warn')
            return
        a, b, tk = self.snap_target(a, b, tk)
        have = [self.has_annot(a, b, 'disfl', d['abbr']) is not None for d in DISFL] + \
               [self.has_annot(a, b, 'error', e['abbr']) is not None for e in ERRS]
        what = '"%s"' % self.events[tk].transcript if tk is not None else '%s - %s' % (fmt_time(a), fmt_time(b))
        dlg = EventPicker(self, have, what)
        if dlg.exec() != QtWidgets.QDialog.Accepted:
            self.set_status('Cancelled - nothing changed.')
            return
        want = dlg.values()
        if want == have:
            return
        self.apply_types(a, b, tk, want)

    def apply_types(self, a, b, tk, want):
        kinds = ['disfl'] * len(DISFL) + ['error'] * len(ERRS)
        cats = [d['abbr'] for d in DISFL] + [e['abbr'] for e in ERRS]
        keep = self.events[tk].uid if tk is not None else 0
        self.push_undo()
        dels = []
        for q in range(len(kinds)):
            j = self.has_annot(a, b, kinds[q], cats[q])
            if j is not None and not want[q]:
                dels.append(j)
        for j in sorted(dels, reverse=True):
            del self.events[j]
        if keep:
            tk = self.find_uid(keep)
        for q in range(len(kinds)):
            if want[q] and self.has_annot(a, b, kinds[q], cats[q]) is None:
                self.add_raw(a, b, tk, kinds[q], cats[q])
        self.relink_all()
        self.cur = None
        self.sel_s, self.sel_e = a, b
        self.events_changed()
        self.update_selection_graphics()
        if self.group_u > 0:
            self.set_status('Marked: %s. Highlight the next disfluent part + Enter, W = whole episode, Esc = done.'
                            % self.piece_types(a, b), 'ok')
        else:
            self.set_status('Marked: %s (Ctrl+Z to undo).' % self.piece_types(a, b), 'ok')

    def word_picker(self, sk, a, b):
        toks = self.events[sk].transcript.split()
        if not toks:
            QtWidgets.QMessageBox.information(self, 'Fluent word', 'That transcript is empty - type it first.')
            return None
        if len(toks) == 1:
            return 1
        names = ['%d.  %s' % (i + 1, t) for i, t in enumerate(toks)]
        frac = ((a + b) / 2 - self.events[sk].start) / max(self.events[sk].end - self.events[sk].start, 1e-9)
        est = min(max(int(math.ceil(frac * len(toks))), 1), len(toks))
        it, ok = QtWidgets.QInputDialog.getItem(self, 'Fluent word', 'Which word is this fluent piece?',
                                                names, est - 1, False)
        return names.index(it) + 1 if ok else None

    def set_part_type(self, k, code):
        ev = self.events[k]
        if (code == 'flu' and ev.kind == 'fluent') or ev.cat == code:
            self.set_status('%s is already that type.' % self.kind_label(k))
            return
        if code == 'flu':
            if self.overlap_idx(ev.start, ev.end, ['fluent'], [k]):
                self.set_status('That overlaps another fluent word.', 'warn')
                return
            sk = self.seg_of(k)
            wi, word = 0, ev.transcript
            if sk is not None:
                wi = self.word_picker(sk, ev.start, ev.end)
                if wi is None:
                    return
                word = norm_word_keep(self.events[sk].transcript.split()[wi - 1])
            self.push_undo()
            ev.kind, ev.cat, ev.link, ev.wi, ev.transcript = 'fluent', '', 0, wi, word
        else:
            self.push_undo()
            if ev.kind == 'fluent':
                ev.transcript, ev.wi = '', 0
            ev.kind, ev.cat = 'disfl', code
        self.relink_all()
        self.events_changed()
        self.set_status('Changed to %s (Ctrl+Z to undo).' % self.kind_label(k), 'ok')

    def change_type_dialog(self):
        if self.cur is None:
            QtWidgets.QMessageBox.information(self, 'Change type', 'Select an item first.')
            return
        k = self.cur
        kind = self.events[k].kind
        if kind == 'trans':
            QtWidgets.QMessageBox.information(self, 'Change type',
                                              'A transcript has no type. Select a fluent word, disfluency or error.')
        elif kind in ('fluent', 'disfl'):
            codes = ['flu'] + [d['abbr'] for d in DISFL]
            names = ['flu - fluent word'] + ['%s - %s' % (d['abbr'], d['name']) for d in DISFL]
            cur = 0 if kind == 'fluent' or self.events[k].cat not in codes else codes.index(self.events[k].cat)
            it, ok = QtWidgets.QInputDialog.getItem(self, 'Change type', 'Choose the type:', names, cur, False)
            if ok:
                self.set_part_type(k, codes[names.index(it)])
        else:
            q = self.error_picker('Change error type')
            if q is not None and self.events[k].cat != ERRS[q]['abbr']:
                self.push_undo()
                self.events[k].cat = ERRS[q]['abbr']
                self.events_changed()
                self.set_status('Changed to %s (Ctrl+Z to undo).' % ERRS[q]['abbr'], 'ok')

    def delete_current_event(self):
        if self.cur is None or self.cur >= len(self.events):
            self.set_status('Nothing selected - click an item first, then Delete.', 'warn')
            return
        self.commit_edit_quiet()
        k = self.cur
        ev = self.events[k]
        if ev.kind == 'trans':
            att = self.attached_to(k)
            if att and all(self.events[j].kind == 'fluent' for j in att):
                self.push_undo()
                self._delete_idx([k] + att)
            elif att:
                mb = QtWidgets.QMessageBox(self)
                mb.setWindowTitle('Delete piece')
                mb.setText('Delete the piece "%s" and what it is marked as (%s)?'
                           % (ev.transcript, self.piece_types(ev.start, ev.end)))
                b_all = mb.addButton('Delete all', QtWidgets.QMessageBox.AcceptRole)
                b_one = mb.addButton('Piece only', QtWidgets.QMessageBox.DestructiveRole)
                mb.addButton('Cancel', QtWidgets.QMessageBox.RejectRole)
                mb.exec()
                if mb.clickedButton() is b_all:
                    self.push_undo()
                    self._delete_idx([k] + att)
                elif mb.clickedButton() is b_one:
                    self.push_undo()
                    for j in att:
                        self.events[j].tok = 0
                    self._delete_idx([k])
                else:
                    return
            else:
                self.push_undo()
                self._delete_idx([k])
        else:
            self.push_undo()
            self._delete_idx([k])
        self.relink_all()
        if self.group_u > 0 and self.find_uid(self.group_u) is None:
            self.group_u = 0
        self.cur = None
        self.sel_s, self.sel_e = NAN, NAN
        self.events_changed()
        self.update_selection_graphics()
        self.set_status('Deleted (Ctrl+Z to undo).', 'ok')

    def _delete_idx(self, idx):
        for j in sorted(set(idx), reverse=True):
            del self.events[j]

    def edit_current_event(self):
        if self.cur is None:
            QtWidgets.QMessageBox.information(self, 'Edit text', 'Select an item first.')
            return
        k = self.cur
        if self.events[k].start < self.viewS or self.events[k].end > self.viewE:
            self.select_event(k, True)
        self.start_editing(k, 'notes' if self.events[k].kind == 'error' else 'text')

    def edit_notes(self):
        if self.cur is None or self.cur >= len(self.events):
            self.set_status('Click an item first, then N to write its notes (or double-click a highlight on the '
                            'Notes strip).', 'warn')
            return
        self.start_editing(self.cur, 'notes')

    def undo(self):
        if not self.undo_stack:
            self.set_status('Nothing to undo.', 'warn')
            return
        self.commit_edit_quiet()
        self.events, self.next_uid = self.undo_stack.pop()
        if self.group_u > 0 and self.find_uid(self.group_u) is None:
            self.group_u = 0
        self.cur = None
        self.sel_s, self.sel_e = NAN, NAN
        self.events_changed()
        self.update_selection_graphics()
        self.set_status('Undid the last change.', 'ok')

    # ======================= text editing =====================================
    def start_editing(self, k, panel):
        if not self.loaded or k is None or k >= len(self.events):
            return
        if self.cur != k:
            self.cur = k
            self.sel_s, self.sel_e = self.events[k].start, self.events[k].end
            self.sync_list_selection()
            self.update_selection_graphics()
        self.redraw_events()
        self.begin_typing(panel)

    def begin_typing(self, panel):
        if not self.loaded or self.cur is None or self.cur >= len(self.events):
            return
        if self.editing:
            self.commit_box()
        k = self.cur
        ev = self.events[k]
        if panel == 'notes' or ev.kind == 'error':
            self.box_kind, fld = 'notes', 'notes'
        else:
            self.box_kind = 'transcript' if ev.kind == 'trans' else 'detext'
            fld = 'transcript'
        self.box.setText(getattr(ev, fld))
        self.box_target = k
        self.editing = True
        self.position_box()
        if not self.box.isVisible():
            self.editing = False
            return
        self.box.setFocus()
        self.box.end(False)
        self.set_status('Typing the %s for %s - Enter saves, Esc cancels.'
                        % ('notes' if fld == 'notes' else 'text', self.event_label(k)))

    def position_box(self):
        if not self.editing:
            self.box.hide()
            return
        vb = {'transcript': self.vb_trans, 'detext': self.vb_de, 'notes': self.vb_notes}[self.box_kind]
        if self.box_target == -1:
            if self.pending_new is None:
                self.box.hide()
                return
            a, b = self.pending_new
        elif self.box_target is not None and self.box_target < len(self.events):
            a, b = self.events[self.box_target].start, self.events[self.box_target].end
        else:
            self.box.hide()
            return
        x0 = min(max(a, self.viewS), self.viewE)
        x1 = max(min(b, self.viewE), self.viewS)
        r = vb.sceneBoundingRect()
        xa = vb.mapViewToScene(QPointF(x0, 0)).x()
        xb = vb.mapViewToScene(QPointF(x1, 0)).x()
        pa = self.glw.mapFromScene(QPointF(xa, r.top()))
        pb = self.glw.mapFromScene(QPointF(xb, r.bottom()))
        left = self.glw.mapFromScene(r.topLeft()).x()
        right = self.glw.mapFromScene(r.topRight()).x()
        w = min(max(420, pb.x() - pa.x()), right - left)
        x = max(left, min(pa.x(), right - w))
        h = 28
        y = int((pa.y() + pb.y()) / 2 - h / 2)
        self.box.setGeometry(int(x), y, int(w), h)
        self.box.show()
        self.box.raise_()

    def commit_box(self):
        if not self.editing:
            return
        s = self.box.text().strip()
        self.box.hide()
        self.editing = False
        target = self.box_target
        self.box_target = None
        self.glw.setFocus()
        if target == -1:
            pn = self.pending_new
            self.pending_new = None
            if pn is None:
                self.redraw_events()
                return
            if not s:
                self.redraw_events()
                self.set_status('Nothing added - it needs text. Double-click the highlight to try again.', 'warn')
                return
            self.push_undo()
            self.add_pieces(pn[0], pn[1], s)
            self.cur = None
            self.sel_s, self.sel_e = NAN, NAN
            self.events_changed()
            self.update_selection_graphics()
            self.set_status('Added "%s". Drag the edges between pieces to fit the audio, then highlight the fluent '
                            'word (e.g. table) and press Enter.' % s, 'ok')
            return
        if target is None or target >= len(self.events):
            self.redraw_events()
            return
        fld = 'notes' if self.box_kind == 'notes' else 'transcript'
        if getattr(self.events[target], fld) != s:
            self.push_undo()
            setattr(self.events[target], fld, s)
            if fld == 'transcript' and self.events[target].kind == 'trans':
                for j in self.attached_to(target):
                    if self.events[j].kind != 'error':
                        self.events[j].transcript = token_word(s)
            self.dirty = True
            self.update_title()
            self.update_event_list()
            self.redraw_events()
            self.set_status('Saved %s for %s.' % (fld, self.event_label(target)), 'ok')
        else:
            self.redraw_events()

    def cancel_typing(self):
        if not self.editing:
            return
        self.editing = False
        self.box.hide()
        self.glw.setFocus()
        if self.box_target == -1:
            self.pending_new = None
            self.set_status('Cancelled - nothing added.')
        else:
            self.set_status('Edit cancelled - text unchanged.')
        self.box_target = None
        self.redraw_events()

    def commit_edit(self):
        if self.editing:
            self.commit_box()

    def commit_edit_quiet(self):
        self.editing = False
        self.pending_new = None
        self.box_target = None
        self.box.hide()

    def insert_symbol(self, sym):
        if self.editing:
            self.box.insert(sym)
            self.box.setFocus()
            self.set_status('Inserted %s - keep typing; Enter saves.' % sym)
        else:
            QtWidgets.QApplication.clipboard().setText(sym)
            self.set_status('%s copied. Start typing in a transcript / notes box first to insert it directly.' % sym)

    # ======================= events table =====================================
    def update_event_list(self):
        t = self.tree
        t.blockSignals(True)
        t.clear()
        E = self.events
        fl = sorted([i for i, e in enumerate(E) if e.kind == 'fluent'], key=lambda i: E[i].start)
        self.expanded &= {E[i].uid for i in fl}
        selU = self.group_u if self.group_u > 0 else self.group_uid(self.cur)
        for f in fl:
            u = E[f].uid
            kids = sorted(self.linked_all(u), key=lambda i: E[i].start)
            lab = E[f].transcript + ('  (%d)' % len(kids) if kids else '')
            it = QtWidgets.QTreeWidgetItem([fmt_time(E[f].start), fmt_time(E[f].end), lab,
                                            one_line(self.fluent_notes(f))])
            it.setData(0, Qt.UserRole, ('fluent', f))
            if u == selU:
                for c in range(4):
                    it.setBackground(c, QBrush(qc((0.80, 0.88, 1.0))))
            t.addTopLevelItem(it)
            for j in kids:
                ch = QtWidgets.QTreeWidgetItem([fmt_time(E[j].start), fmt_time(E[j].end),
                                                '%s  %s' % (self.event_code(j), E[j].cat), one_line(E[j].notes)])
                ch.setData(0, Qt.UserRole, ('child', j))
                for c in range(4):
                    ch.setBackground(c, QBrush(qc(mix(self.event_color(j), 0.45, 0.55))))
                it.addChild(ch)
            it.setExpanded(u in self.expanded)
        orph = sorted([i for i, e in enumerate(E) if e.kind in ('disfl', 'error')
                       and (e.link <= 0 or self.find_uid(e.link) is None)], key=lambda i: E[i].start)
        if orph:
            it = QtWidgets.QTreeWidgetItem(['', '', '(unresolved)  (%d)' % len(orph), ''])
            it.setData(0, Qt.UserRole, ('unresolved', -1))
            for c in range(4):
                it.setBackground(c, QBrush(qc((0.92, 0.92, 0.92))))
            t.addTopLevelItem(it)
            for j in orph:
                ch = QtWidgets.QTreeWidgetItem([fmt_time(E[j].start), fmt_time(E[j].end),
                                                '%s  %s' % (self.event_code(j), E[j].cat), one_line(E[j].notes)])
                ch.setData(0, Qt.UserRole, ('child', j))
                for c in range(4):
                    ch.setBackground(c, QBrush(qc(mix(self.event_color(j), 0.45, 0.55))))
                it.addChild(ch)
            it.setExpanded(self.unresolved_open)
        if not fl and not orph:
            it = QtWidgets.QTreeWidgetItem(['', '', '(no fluent events yet)', ''])
            it.setData(0, Qt.UserRole, ('header', -1))
            t.addTopLevelItem(it)
        t.blockSignals(False)

    def on_tree_expand(self, item, open_):
        kind, k = item.data(0, Qt.UserRole)
        if kind == 'unresolved':
            self.unresolved_open = open_
        elif kind == 'fluent' and 0 <= k < len(self.events):
            u = self.events[k].uid
            if open_:
                self.expanded.add(u)
            else:
                self.expanded.discard(u)

    def on_tree_click(self, item, col):
        data = item.data(0, Qt.UserRole)
        QTimer.singleShot(0, lambda: self.cb(lambda: self._tree_click(*data)))

    def _tree_click(self, kind, k):
        if not self.loaded:
            return
        self.commit_edit()
        if kind == 'unresolved':
            self.unresolved_open = not self.unresolved_open
            self.update_event_list()
            return
        if k < 0 or k >= len(self.events):
            return
        if kind == 'fluent':
            u = self.events[k].uid
            if u in self.expanded:
                self.expanded.discard(u)
            else:
                self.expanded.add(u)
            t = self.find_uid(self.events[k].tok)
            self.select_event(k if t is None else t, True)
            self.set_status('Fluent event "%s" - its disfluency / error events are listed under it. Space = play, '
                            'Enter = add more parts to this word.' % self.events[k].transcript)
        else:
            self.select_event(k, True)
            self.set_status('Selected %s. Enter = change its type(s), N = notes, Delete = remove.'
                            % self.kind_label(k))

    def sync_list_selection(self):
        self.update_event_list()

    def select_event(self, k, ensure_visible):
        self.cur = k
        self.sel_s, self.sel_e = self.events[k].start, self.events[k].end
        self.cursor = self.sel_s
        if ensure_visible and (self.sel_s < self.viewS or self.sel_e > self.viewE):
            w = max(self.viewE - self.viewS, 1.5 * (self.sel_e - self.sel_s))
            c = (self.sel_s + self.sel_e) / 2
            self.set_view(c - w / 2, c + w / 2)
        else:
            self.redraw_events()
            self.update_selection_graphics()
            self.update_cursor_graphics()
        self.show_frame_at(self.cursor)
        self.sync_list_selection()

    # ======================= playback =========================================
    def toggle_play(self):
        if self.playing:
            self.pause_play()
        else:
            self.start_play()

    def set_loop(self, on, quiet=False):
        self.loop_on = bool(on)
        if self.chk_loop.isChecked() != self.loop_on:
            self.chk_loop.blockSignals(True)
            self.chk_loop.setChecked(self.loop_on)
            self.chk_loop.blockSignals(False)
        if self.aeng is not None:
            self.aeng.loop_on = self.loop_on
        self.set_status('Loop on: the selection replays until you pause.' if self.loop_on else
                        'Loop off: playback stops at the end of the selection.')

    def set_speed_idx(self, i):
        if not 1 <= i <= len(SPEEDS):
            return
        self.play_speed = SPEEDS[i - 1]
        self.pop_speed.setCurrentIndex(i - 1)
        if self.playing:
            t = self.aeng.current_time()
            self.stop_play()
            self.cursor = self.clamp_t(t)
            self.start_play()
        msg = 'Speed %s.' % SPEED_NAMES[i - 1]
        self.set_status(msg)

    def start_play(self):
        if not self.loaded or self.playing:
            return
        if self.has_sel():
            a, b = self.sel_s, self.sel_e
            if not isnan(self.cursor) and a < self.cursor < b - 0.01:
                t0, what = self.cursor, 'rest of selection'
            else:
                t0, what = a, 'selection'
            t1, loop_a, cr = b, a, (a, b)
            if a < self.viewS or b > self.viewE:
                w = max(self.viewE - self.viewS, 1.2 * (b - a))
                c = (a + b) / 2
                self.set_view(c - w / 2, c + w / 2)
        elif not isnan(self.cursor):
            t0, t1, loop_a, cr, what = self.cursor, self.dur, None, None, 'from cursor'
        else:
            t0, t1, loop_a, cr, what = self.viewS, self.dur, None, None, 'from view start'
        if t1 - t0 <= 1.0 / self.fs:
            return
        self.play_loop_a = loop_a
        self.aeng.loop_on = self.loop_on
        if self.play_speed != 1 and t1 - t0 > 5:
            self.set_status('Preparing slowed-down audio...')
            QtWidgets.QApplication.processEvents()
        self.play_end = self.aeng.play(t0, t1, self.play_speed, loop_a)
        self.worker.begin(t0, cr)
        self.playing = True
        for p in self.play_lines:
            p.setPos(t0)
            p.show()
        self.play_timer.start()
        self.sync_play_btn()
        msg = 'Playing %s at %.2fx...   Space = pause.' % (what, self.play_speed)
        if self.loop_on and loop_a is not None:
            msg += '  Looping (L turns it off).'
        if self.aeng.silent:
            msg += '  (No audio device - playing silently.)'
        s0, s1 = int(t0 * self.fs), int(t1 * self.fs)
        if np.max(np.abs(self.audio[s0:s1])) < 1e-4:
            msg += '  (This part of the recording is silent.)'
        self.set_status(msg)

    def play_tick(self):
        if not self.playing:
            return
        t = self.aeng.current_time()
        looping = self.loop_on and self.play_loop_a is not None
        if self.aeng.finished or (not looping and t >= self.play_end - 1e-3):
            self.stop_play()
            return
        if self.play_loop_a is None and (t > self.viewE or t < self.viewS):
            w = self.viewE - self.viewS
            self.viewS = max(0.0, t - 0.02 * w)
            self.viewE = min(self.dur, self.viewS + w)
            self.refresh_view()
        for p in self.play_lines:
            p.setPos(t)
        self.ov_cur.setPos(t)
        self.worker.set_target(t)
        p = self.worker.get(t)
        if p is not None and p is not self.last_shown:
            self.video.set_prepared(p)
            self.last_shown = p

    def stop_play(self):
        self.play_timer.stop()
        if self.aeng is not None:
            self.aeng.stop()
        if self.worker is not None:
            self.worker.pause()
        self.playing = False
        for p in self.play_lines:
            p.hide()
        self.ov_cur.setPos(self.cursor if not isnan(self.cursor) else -1)
        self.sync_play_btn()

    def pause_play(self):
        if not self.playing:
            return
        t = self.aeng.current_time()
        drops = self.aeng.underruns
        self.stop_play()
        self.cursor = self.clamp_t(t)
        self.update_cursor_graphics()
        self.show_frame_at(self.cursor)
        msg = 'Paused at %s. Press Space to continue from here.' % fmt_time(self.cursor)
        if drops:
            msg += '  [audio dropouts: %d]' % drops
        self.set_status(msg, 'warn' if drops else 'info')

    def start_play_fresh(self):
        self.stop_play()
        a, b, _ = self.current_target()
        if not isnan(a):
            self.sel_s, self.sel_e = a, b
        self.cursor = self.sel_s
        self.update_selection_graphics()
        self.start_play()

    def play_current_event(self):
        if self.cur is None:
            QtWidgets.QMessageBox.information(self, 'Play event', 'Select an event first.')
            return
        self.stop_play()
        self.select_event(self.cur, True)
        self.start_play()

    # ======================= video ============================================
    def show_frame_at(self, t):
        if not self.loaded or isnan(t):
            return
        p = self.worker.cached(t) if self.worker is not None else None
        if p is not None:
            self.video.set_prepared(p)
            self.last_shown = p
            return
        try:
            x = self.vsrc.frame_at(t)
        except Exception:
            x = None
        if x is None:
            return
        self.last_raw = frame_rgb(x[0])
        self.last_raw_t = x[1]
        self.video.set_prepared(prepare(self.last_raw, x[1], self.vparams.snapshot()))

    def schedule_video_params(self):
        if self.loaded:
            self.vp_timer.start()

    def apply_video_params(self, quiet=True):
        if not self.loaded:
            return
        dpr = self.video.devicePixelRatioF()
        self.vparams.update(tw=int(self.video.width() * dpr), th=int(self.video.height() * dpr))
        if self.worker is not None:
            self.worker.invalidate()
        if self.last_raw is not None:
            self.video.set_prepared(prepare(self.last_raw, self.last_raw_t, self.vparams.snapshot()))

    def set_video_zoom(self, z, quiet=False):
        self.vparams.update(zoom=z)
        self.apply_video_params()
        if quiet:
            return
        if z is None:
            self.set_status('Video zoom reset (whole picture).')
        else:
            W, _ = self.vparams.dims()
            self.set_status('Video zoom %.1fx - scroll on the video to zoom, drag it to move around, '
                            'double-click or Fit to reset.' % (W / z[2]))

    def video_zoom_by(self, f, cx=None, cy=None):
        if not self.loaded:
            return
        W, H = self.vparams.dims()
        z = self.vparams.zoom or (0, 0, W, H)
        if cx is None:
            cx, cy = z[0] + z[2] / 2, z[1] + z[3] / 2
        w = z[2] / f
        if w >= W - 1:
            self.set_video_zoom(None)
            return
        w = max(int(round(w)), 40)
        h = max(int(round(w * H / W)), 20)
        x0 = int(round(min(max(cx - w / 2, 0), W - w)))
        y0 = int(round(min(max(cy - h / 2, 0), H - h)))
        self.set_video_zoom((x0, y0, w, h))

    def reset_video_zoom(self):
        self.set_video_zoom(None)

    def rotate_video(self):
        rot = (self.vparams.rot + 3) % 4      # 90 degrees clockwise each click
        self.vparams.update(rot=rot, zoom=None)
        self.settings.setValue('rot/' + self.video_path, rot)
        self.apply_video_params()
        self.set_status('Video rotated (remembered for this file).')

    def toggle_pop(self):
        if self.pop is not None:
            self.dock_video()
        else:
            self.pop_out_video()

    def pop_out_video(self):
        if not self.loaded:
            return
        self.pop = QtWidgets.QWidget(None, Qt.Window)
        self.pop.setWindowTitle(APPNAME + ' - video')
        self.pop.setStyleSheet('background:black;')
        lay = QtWidgets.QVBoxLayout(self.pop)
        lay.setContentsMargins(0, 0, 0, 0)
        bar = QtWidgets.QHBoxLayout()
        for lab, fn in (('+', lambda: self.video_zoom_by(1.5)), ('-', lambda: self.video_zoom_by(1 / 1.5)),
                        ('Fit', self.reset_video_zoom), ('Dock', self.dock_video)):
            b = QtWidgets.QPushButton(lab)
            b.setFocusPolicy(Qt.NoFocus)
            b.setStyleSheet('background:#ddd;color:black;')
            b.clicked.connect(lambda _=False, f=fn: self.cb(f))
            bar.addWidget(b)
        hint = QtWidgets.QLabel('Scroll = zoom  |  drag = move  |  double-click = fit  |  Space = play')
        hint.setStyleSheet('color:#ddd;')
        bar.addWidget(hint, 1)
        lay.addLayout(bar)
        self.vid_col.removeWidget(self.video)
        lay.addWidget(self.video, 1)
        self.dock_msg.show()
        self.btn_pop.setText('Dock video')
        W, H = self.vparams.dims()
        scr = QtWidgets.QApplication.primaryScreen().availableGeometry()
        w = min(1100, scr.width() - 80)
        h = min(int(w * H / W) + 40, scr.height() - 120)
        self.pop.resize(w, h)
        self.pop.closeEvent = lambda e: (e.ignore(), self.cb(self.dock_video))
        self.pop.show()
        self.set_status('Video popped out - resize that window for a bigger, sharper picture.', 'ok')
        self.schedule_video_params()

    def dock_video(self):
        if self.pop is None:
            return
        self.pop.layout().removeWidget(self.video)
        self.vid_col.insertWidget(1, self.video, 1)
        self.video.show()
        self.dock_msg.hide()
        self.pop.closeEvent = lambda e: e.accept()
        self.pop.close()
        self.pop.deleteLater()
        self.pop = None
        self.btn_pop.setText('Pop out video')
        self.set_status('Video docked back in the main window.')
        self.schedule_video_params()

    # ======================= loading video ====================================
    def ui_load_video(self):
        if not self.confirm_unsaved('loading another video'):
            return
        fn, _ = QtWidgets.QFileDialog.getOpenFileName(
            self, 'Select a video file', '',
            'Video files (*.mp4 *.mov *.MOV *.MP4 *.avi *.m4v *.mkv *.wmv *.webm);;All files (*)')
        if fn:
            self.load_video(fn)

    def load_video(self, path):
        path = os.path.abspath(os.path.expanduser(path))
        prog = QtWidgets.QProgressDialog('Extracting the audio track...', None, 0, 0, self)
        prog.setWindowTitle('Loading video')
        prog.setMinimumDuration(0)
        prog.show()
        QtWidgets.QApplication.processEvents()
        try:
            audio, fs, origin = load_audio(path)
            vs = VideoSource(path, origin)
            x = vs.frame_at(0.0)
            if x is None:
                raise RuntimeError('No video frames could be decoded.')
        except Exception as err:
            prog.close()
            QtWidgets.QMessageBox.critical(self, 'Load error', 'Could not open the video:\n\n%s' % err)
            return
        prog.close()
        self.stop_play()
        self.commit_edit_quiet()
        if self.worker is not None:
            self.worker.stop()
        if self.aeng is not None:
            self.aeng.close()
        if self.vsrc is not None:
            self.vsrc.close()
        self.vsrc = vs
        self.video_path = path
        self.audio, self.fs = audio, fs
        self.amax = max(float(np.max(np.abs(audio))), 1e-9)
        self.dur = len(audio) / fs
        self.max_freq = min(5000, fs / 2)
        self.vbs['spec'].yr = (0, self.max_freq)
        self.vb_spec.setYRange(0, self.max_freq, padding=0)
        self.events = []
        self.next_uid = 1
        self.cur = None
        self.sel_s, self.sel_e = NAN, NAN
        self.cursor = 0.0
        self.pending_new = None
        self.group_u = 0
        self.undo_stack = []
        self.dirty = False
        self.expanded = set()
        self.viewS, self.viewE = 0.0, min(self.dur, 10.0)
        rot = frame_rotation(x[0])
        saved = self.settings.value('rot/' + path)
        if saved is not None:
            rot = int(saved) % 4
        self.vparams = ViewParams(vs.W, vs.H, rot)
        self.aeng = AudioEngine(audio, fs)
        self.worker = PlaybackWorker(path, origin, self.vparams, vs.fps)
        self.worker.start()
        self.loaded = True
        self.set_video_dependent(True)
        dpr = self.video.devicePixelRatioF()
        self.vparams.update(tw=int(self.video.width() * dpr), th=int(self.video.height() * dpr))
        self.last_raw = frame_rgb(x[0])
        self.last_raw_t = x[1]
        self.video.set_prepared(prepare(self.last_raw, x[1], self.vparams.snapshot()))
        self.build_overview()
        self.update_event_list()
        self.update_title()
        self.refresh_view()
        self.glw.setFocus()
        nm = os.path.basename(path)
        codec = vs.s.codec_context.name
        msg = ('Loaded %s  (%s, %dx%d %s, %.0f fps, %d Hz audio).  Drag to select speech, then double-click it '
               '(or Enter) to type the transcript.' % (nm, fmt_time(self.dur), vs.W, vs.H, codec, vs.fps, fs))
        if sd is None:
            msg += '  [No audio output: install PortAudio / sounddevice for sound.]'
        self.set_status(msg, 'ok')

    # ======================= save / load ======================================
    def build_save_cells(self):
        hdr = ['start', 'stop', 'start-audiofile-relative', 'stop-audiofile-relative', 'event-type',
               'transcript', 'fluent-event-id', 'disfluency', 'error', 'notes']
        if not self.events:
            return [hdr]
        ev = sorted(self.events, key=lambda e: e.start)
        keep = []
        for e in ev:
            if e.kind == 'trans' and any(f.kind == 'fluent' and f.tok == e.uid and abs(f.start - e.start) < 1e-3
                                         and abs(f.end - e.end) < 1e-3 for f in ev):
                continue
            keep.append(e)
        fuids = [e.uid for e in keep if e.kind == 'fluent']
        C = [hdr]
        for e in keep:
            tr, fid, dc, ec = '', '', '', ''
            if e.kind == 'trans':
                ty, tr = 'transcript', e.transcript
            elif e.kind == 'fluent':
                ty, tr, fid = 'fluent', e.transcript, fuids.index(e.uid) + 1
            elif e.kind == 'disfl':
                ty, dc = 'disfluency', e.cat
                if DISFL_HAS_TRANSCRIPT:
                    tr = e.transcript
            else:
                ty, ec = 'error', e.cat
            if e.kind in ('disfl', 'error') and e.link > 0 and e.link in fuids:
                fid = fuids.index(e.link) + 1
            C.append(['' if isnan(e.gStart) else e.gStart, '' if isnan(e.gStop) else e.gStop,
                      e.start, e.end, ty, tr, fid, dc, ec, e.notes])
        return C

    def code_table(self):
        C = [['code', 'label', 'meaning', 'group'], ['F', 'fluent', 'fluent-accurate word (resolution)', 'fluent']]
        for d in DISFL:
            grp = 'typical disfluency (TD)' if d['abbr'] == 'td' else 'stuttering-like disfluency (SLD)'
            C.append([d['key'].upper(), d['abbr'], d['name'], grp])
        for e in ERRS:
            C.append(['E-' + re.sub(r'^er-', '', e['abbr']).upper(), e['abbr'], e['name'], 'error'])
        return C

    def export_cells(self):
        C = [['start', 'end', 'duration', 'fluent_event_id', 'fluent_event', 'event_code', 'notes']]
        E = self.events
        r3 = lambda x: round(x * 1000) / 1000
        fl = sorted([i for i, e in enumerate(E) if e.kind == 'fluent'], key=lambda i: E[i].start)
        for n, f in enumerate(fl, 1):
            C.append([r3(E[f].start), r3(E[f].end), r3(E[f].end - E[f].start), n, E[f].transcript, 'F',
                      self.fluent_notes(f)])
            for j in sorted(self.linked_all(E[f].uid), key=lambda i: E[i].start):
                C.append([r3(E[j].start), r3(E[j].end), r3(E[j].end - E[j].start), n, E[f].transcript,
                          self.event_code(j), E[j].notes])
        orph = sorted([i for i, e in enumerate(E) if e.kind in ('disfl', 'error')
                       and (e.link <= 0 or self.find_uid(e.link) is None)], key=lambda i: E[i].start)
        for j in orph:
            C.append([r3(E[j].start), r3(E[j].end), r3(E[j].end - E[j].start), '', '(unresolved)',
                      self.event_code(j), E[j].notes])
        return C

    def save_annotations(self):
        if not self.loaded:
            QtWidgets.QMessageBox.warning(self, 'Save', 'Load a video first.')
            return False
        self.commit_edit()
        if not self.events:
            r = QtWidgets.QMessageBox.question(self, 'Nothing to save',
                                               'There is nothing to save yet. Save an empty annotation file anyway?')
            if r != QtWidgets.QMessageBox.Yes:
                return False
        d, base = os.path.split(self.video_path)
        default = os.path.join(d, os.path.splitext(base)[0] + '_annot-disfluencies.xlsx')
        out, flt = QtWidgets.QFileDialog.getSaveFileName(self, 'Save annotations', default,
                                                         'Excel workbook (*.xlsx);;CSV files (*.csv)')
        if not out:
            return False
        E, K, A = self.export_cells(), self.code_table(), self.build_save_cells()
        csv_mode = out.lower().endswith('.csv')
        if not csv_mode and not out.lower().endswith('.xlsx'):
            out += '.xlsx'
        try:
            if csv_mode:
                ob = os.path.splitext(out)[0]
                write_csv(out, E)
                write_csv(ob + '_codes.csv', K)
                write_csv(ob + '_annotator.csv', A)
            else:
                write_xlsx(out, [('events', E), ('codes', K), ('annotator_data', A)])
        except Exception as err:
            QtWidgets.QMessageBox.critical(self, 'Save error', 'Could not save the annotations:\n\n%s\n\nIf the '
                                           'file is open in Excel, close it and try again.' % err)
            return False
        nF = sum(e.kind == 'fluent' for e in self.events)
        nD = sum(e.kind in ('disfl', 'error') for e in self.events)
        self.dirty = False
        self.update_title()
        self.set_status('Saved %d fluent event(s) and %d disfluency / error event(s) to %s'
                        % (nF, nD, os.path.basename(out)), 'ok')
        return True

    def load_annotations(self):
        if not self.loaded:
            QtWidgets.QMessageBox.warning(self, 'Load annotations',
                                          'Load the matching video first, then load its annotations.')
            return
        self.commit_edit()
        f, _ = QtWidgets.QFileDialog.getOpenFileName(self, 'Load annotations', os.path.dirname(self.video_path),
                                                     'Annotation tables (*.xlsx *.csv)')
        if not f:
            return
        fn = os.path.basename(f)
        try:
            C = read_saved_table(f)
        except Exception as err:
            QtWidgets.QMessageBox.critical(self, 'Load error', 'Could not read "%s".\n\n%s\n\nTry re-saving it from '
                                           'Excel as .xlsx or .csv and load it again.' % (fn, err))
            return
        result = self.parse_annotation_table(C, fn)
        if result is None:
            return
        new_ev, n_new = result
        merge = False
        if self.events:
            mb = QtWidgets.QMessageBox(self)
            mb.setWindowTitle('Load annotations')
            mb.setText('"%s" has %d item(s). This video already has %d.\n\nMerge them into your current annotations, '
                       'or replace the current ones? Ctrl+Z undoes either.' % (fn, len(new_ev), len(self.events)))
            bm = mb.addButton('Merge into current', QtWidgets.QMessageBox.AcceptRole)
            br = mb.addButton('Replace current', QtWidgets.QMessageBox.DestructiveRole)
            mb.addButton('Cancel', QtWidgets.QMessageBox.RejectRole)
            mb.exec()
            if mb.clickedButton() is bm:
                merge = True
                dup, dup_uid = self.compare_to_current(new_ev)
                for i in [i for i, d in enumerate(dup) if d]:
                    for e in new_ev:
                        if e.link == new_ev[i].uid:
                            e.link = dup_uid[i]
                new_ev = [e for e, d in zip(new_ev, dup) if not d]
                if not new_ev:
                    QtWidgets.QMessageBox.information(self, 'Nothing new',
                                                      'Everything in that file is already on the timeline.')
                    return
            elif mb.clickedButton() is not br:
                return
        self.commit_edit_quiet()
        self.push_undo()
        self.group_u = 0
        if merge:
            self.events = self.events + new_ev
            self.dirty = True
            verb = 'Merged'
        else:
            self.events = new_ev
            self.dirty = False
            verb = 'Loaded'
        self.cur = None
        self.sel_s, self.sel_e = NAN, NAN
        self.update_event_list()
        self.update_title()
        self.refresh_view()
        self.set_status('%s %d item(s) from %s.' % (verb, len(new_ev), fn), 'ok')

    def parse_annotation_table(self, C, fn):
        if len(C) < 2:
            QtWidgets.QMessageBox.warning(self, 'Nothing to load', '"%s" has no annotation rows.' % fn)
            return None
        hdrN = [re.sub(r'[^a-z0-9]', '', to_str(h).lower()) for h in C[0]]
        col = lambda nm: hdrN.index(nm) + 1 if nm in hdrN else 0
        cS, cE = col('startaudiofilerelative'), col('stopaudiofilerelative')
        cTyp, cTr, cWi, cId = col('eventtype'), col('transcript'), col('wordindex'), col('fluenteventid')
        cD, cEr, cN, cGS, cGE, cOld = col('disfluency'), col('error'), col('notes'), col('start'), col('stop'), 0
        if cS and cE:
            fmt = 'typed' if cTyp else 'v3'
        else:
            cS, cE, cOld, cTyp, cGS, cGE, fmt = col('starts'), col('ends'), col('eventtype'), 0, 0, 0, 'v2'
            if not cS or not cE:
                QtWidgets.QMessageBox.warning(self, 'Invalid file',
                                              'This file does not look like an annotation file.\n\nExpected columns: '
                                              'start-audiofile-relative, stop-audiofile-relative, event-type, '
                                              'transcript, fluent-event-id, disfluency, error, notes.')
                return None

        def cell(row, c):
            return row[c - 1] if c and c - 1 < len(row) else ''

        R = []
        n_bad = n_out = 0
        for i in range(1, len(C)):
            row = C[i]
            if all(to_str(v) == '' for v in row):
                continue
            st, en = to_num(cell(row, cS)), to_num(cell(row, cE))
            if isnan(st) or isnan(en) or en <= st:
                n_bad += 1
                continue
            if st >= self.dur:
                n_out += 1
                continue
            dc, ec = to_str(cell(row, cD)).lower(), to_str(cell(row, cEr)).lower()
            if fmt == 'typed':
                typ = to_str(cell(row, cTyp)).lower()
            elif fmt == 'v3':
                typ = 'both' if (dc and ec) else ('disfluency' if dc else ('error' if ec else 'fluent'))
            else:
                typ = 'disfluency'
                old = to_str(cell(row, cOld)).lower()
                if old == 'repetition':
                    dc = 'rep'
                elif old == 'block':
                    dc = 'bl'
                elif old == 'prolongation':
                    dc = 'prol'
                elif old in ('text-only', ''):
                    typ = 'transcript'
                else:
                    dc = old
            R.append(SimpleNamespace(row=i + 1, st=st, en=min(en, self.dur), typ=typ, tr=to_str(cell(row, cTr)),
                                     fid=to_num(cell(row, cId)), wi=to_num(cell(row, cWi)), dc=dc, ec=ec,
                                     nt=to_str(cell(row, cN)), gs=to_num(cell(row, cGS)), ge=to_num(cell(row, cGE))))
        if not R:
            QtWidgets.QMessageBox.warning(self, 'Nothing to load', 'No usable rows were found in "%s".' % fn)
            return None

        rows = lambda mask: ', '.join(str(r.row) for r, m in zip(R, mask) if m)[:200]
        is_di = [r.typ in ('disfluency', 'both') for r in R]
        is_er = [r.typ in ('error', 'both') for r in R]
        is_fl = [r.typ == 'fluent' for r in R]
        is_tx = [r.typ == 'transcript' for r in R]
        is_tg = [r.typ == 'target' for r in R]
        has_tr = [bool(r.tr) for r in R]
        issues = []
        if n_bad:
            issues.append('- %d row(s) with missing or invalid times were skipped.' % n_bad)
        if n_out:
            issues.append('- %d row(s) start after the end of this video and were skipped.' % n_out)
        if fmt == 'v2':
            issues.append('- Old layout (starts / ends / event_type): text-only rows load as transcript, the rest as '
                          'unresolved disfluencies.')
        if any(R[i + 1].st < R[i].st for i in range(len(R) - 1)):
            issues.append('- Rows are not sorted by start-audiofile-relative.')
        if fmt == 'typed':
            b = [r.typ not in ('transcript', 'fluent', 'disfluency', 'error', 'target') for r in R]
            if any(b):
                issues.append('- Unknown event-type (rows %s); skipped.' % rows(b))
            b = [(t or f) and (r.dc or r.ec) for r, t, f in zip(R, is_tx, is_fl)]
            if any(b):
                issues.append('- Transcript / fluent rows have a disfluency / error label (rows %s) - ignored.' % rows(b))
        b = [e and not d and h for e, d, h in zip(is_er, is_di, has_tr)]
        if any(b):
            issues.append('- Error rows have a transcript (rows %s) - it will be dropped.' % rows(b))
        if not DISFL_HAS_TRANSCRIPT:
            b = [d and h for d, h in zip(is_di, has_tr)]
            if any(b):
                issues.append('- Disfluency rows have a transcript (rows %s).' % rows(b))
        b = [d and r.dc not in [x['abbr'] for x in DISFL] for r, d in zip(R, is_di)]
        if any(b):
            issues.append('- Unknown disfluency labels (rows %s).' % rows(b))
        b = [e and r.ec not in [x['abbr'] for x in ERRS] for r, e in zip(R, is_er)]
        if any(b):
            issues.append('- Unknown error labels (rows %s).' % rows(b))
        b = [f and not h for f, h in zip(is_fl, has_tr)]
        if any(b):
            issues.append('- Fluent rows with no word (rows %s).' % rows(b))
        fl_ids = [r.fid for r, f in zip(R, is_fl) if f]
        if fmt in ('typed', 'v3'):
            if any(isnan(v) for v in fl_ids):
                issues.append('- Fluent rows without a fluent-event-id.')
            v = [x for x in fl_ids if not isnan(x)]
            if len(set(v)) < len(v):
                issues.append('- Fluent-event-ids are repeated across fluent rows.')
            elif v != list(range(1, len(fl_ids) + 1)):
                issues.append('- Fluent-event-ids are not sequential integers from 1 (renumbered on save).')
            b = [(d or e) and not isnan(r.fid) and r.fid not in fl_ids for r, d, e in zip(R, is_di, is_er)]
            if any(b):
                issues.append('- Disfluency / error rows point to a missing fluent-event-id (rows %s) - loaded '
                              'unresolved.' % rows(b))
        if issues:
            r = QtWidgets.QMessageBox.question(self, 'Annotation file check', 'Checks found problems in "%s":\n\n%s'
                                               '\n\nLoad it anyway?' % (fn, '\n'.join(issues)))
            if r != QtWidgets.QMessageBox.Yes:
                self.set_status('Load cancelled.')
                return None

        new = []

        def mk(rr, kind, cat, tr, link):
            e = Ev(rr.st, rr.en, kind, cat, tr, rr.nt, self.next_uid, link, 0, 0, rr.gs, rr.ge)
            self.next_uid += 1
            new.append(e)
            return e

        tx_spans = [(r.st, r.en) for r, t in zip(R, is_tx) if t]
        for r, t in zip(R, is_tx):
            if t:
                mk(r, 'trans', '', r.tr, 0)
        for r, t in zip(R, is_tg):
            if t and not any(a < r.en and b > r.st for a, b in tx_spans):
                mk(r, 'trans', '', r.tr, 0)
        id_map = {}
        for r, f in zip(R, is_fl):
            if not f:
                continue
            if not any(a <= r.st + 1e-6 and b >= r.en - 1e-6 for a, b in tx_spans):
                mk(r, 'trans', '', r.tr, 0)
            e = mk(r, 'fluent', '', norm_word_keep(r.tr), 0)
            e.wi = self._pick_word_idx(new, e, r.wi)
            if not isnan(r.fid) and r.fid not in id_map:
                id_map[r.fid] = e.uid
        for r, d, er in zip(R, is_di, is_er):
            if not (d or er):
                continue
            lk = id_map.get(r.fid, 0) if (fmt != 'v2' and not isnan(r.fid)) else 0
            if d:
                mk(r, 'disfl', r.dc, r.tr if DISFL_HAS_TRANSCRIPT else '', lk)
            if er:
                mk(r, 'error', r.ec, '', lk)
        trs = [e for e in new if e.kind == 'trans']
        for e in new:
            if e.kind == 'trans':
                continue
            m = next((t for t in trs if abs(t.start - e.start) < 1e-3 and abs(t.end - e.end) < 1e-3), None)
            if m is not None:
                e.tok = m.uid
        return new, len(new)

    def _pick_word_idx(self, new, fe, given):
        sk = next((e for e in new if e.kind == 'trans' and e.start <= fe.start + 1e-6 and e.end >= fe.end - 1e-6),
                  None)
        if sk is None:
            return 0
        toks = [norm_word(t) for t in sk.transcript.split()]
        w = norm_word(fe.transcript)
        if not isnan(given) and 1 <= given <= len(toks) and toks[int(given) - 1] == w:
            return int(given)
        taken = {g.wi for g in new if g.kind == 'fluent' and g is not fe and g.start >= sk.start - 1e-6
                 and g.end <= sk.end + 1e-6}
        for i, t in enumerate(toks, 1):
            if t == w and i not in taken:
                return i
        return 0

    def compare_to_current(self, new_ev):
        dup, uid = [], []
        for n in new_ev:
            d = next((e for e in self.events if abs(e.start - n.start) < 1e-3 and abs(e.end - n.end) < 1e-3
                      and e.kind == n.kind and e.cat == n.cat), None)
            dup.append(d is not None)
            uid.append(d.uid if d is not None else 0)
        return dup, uid

    def confirm_unsaved(self, action):
        if not self.dirty:
            return True
        mb = QtWidgets.QMessageBox(self)
        mb.setWindowTitle('Unsaved changes')
        mb.setText('You have unsaved annotation changes.\n\nSave them before %s?' % action)
        bs = mb.addButton('Save', QtWidgets.QMessageBox.AcceptRole)
        bd = mb.addButton("Don't save", QtWidgets.QMessageBox.DestructiveRole)
        mb.addButton('Cancel', QtWidgets.QMessageBox.RejectRole)
        mb.exec()
        if mb.clickedButton() is bs:
            return self.save_annotations()
        return mb.clickedButton() is bd

    # ======================= misc =============================================
    def set_status(self, msg, kind='info'):
        col = {'warn': '#bf5900', 'ok': '#00731a'}.get(kind, '#262626')
        self.status.setStyleSheet('color:%s;' % col)
        self.status.setText(msg)

    def update_title(self):
        if not self.video_path:
            self.setWindowTitle(APPNAME)
            return
        self.setWindowTitle('%s - %s%s' % (APPNAME, os.path.basename(self.video_path),
                                           '  * unsaved changes' if self.dirty else ''))

    def show_welcome(self):
        if self.settings.value('hideWelcome', False, type=bool):
            return
        mb = QtWidgets.QMessageBox(self)
        mb.setWindowTitle('Welcome to ' + APPNAME)
        mb.setText(
            'How it works (example "the t-t-table is on the ground"):\n\n'
            '1. Highlight the speech, DOUBLE-CLICK it, type the sentence, Enter.\n'
            '   It appears on the Transcript strip as pieces: the  t-  t-  table ...\n'
            '   (Double-click a highlight on the Notes strip to type notes.)\n'
            '2. Highlight "table" (or click its piece) and press Enter:\n'
            '   FLUENT EVENT SELECTED.\n'
            '3. Highlight "t-t" (or one "t-", or part of a word) and press Enter.\n'
            '   Tick the type(s): rep / prol / bl / td and / or errors (keys\n'
            '   R P B T, 1-6). Enter to finish that part. Repeat for other parts.\n'
            '   W = the whole episode (t-t-table) at once.\n'
            '4. Esc when you are done with that word.\n\n'
            'Space plays the highlight (L loop, 1-4 speed). N = notes. Ctrl/Cmd+S saves.\n'
            'F1 or Help = full instructions.')
        mb.addButton('Get started', QtWidgets.QMessageBox.AcceptRole)
        bh = mb.addButton("Don't show again", QtWidgets.QMessageBox.RejectRole)
        mb.exec()
        if mb.clickedButton() is bh:
            self.settings.setValue('hideWelcome', True)

    def show_help(self):
        msg = (
            'THE WORKFLOW  (example "the t-t-table is on the ground")\n'
            '  1. TRANSCRIPT: highlight the speech on the spectrogram, DOUBLE-CLICK the\n'
            '     highlight, type   the t-t-table is on the ground   and press Enter.\n'
            '     It is split into pieces on the Transcript strip (drag the edges\n'
            '     between pieces to fit).\n'
            '  2. FLUENT WORD: highlight "table" (or click its piece) and press Enter.\n'
            '  3. DISFLUENT PARTS: highlight "t-t" (or one "t-") and press Enter.\n'
            '     Tick one or more:  R rep   P prol   B bl   T td   /   1-6 errors,\n'
            '     then Enter. W = the whole episode.\n'
            '  4. Esc = done with this word.\n'
            '  Shortcuts: R / P / B / T / E add a type straight away. Enter on a\n'
            '  selected disfluency changes its types.\n\n'
            'NOTES\n'
            '  Double-click a highlight on the Notes strip, or click an item and press N.\n\n'
            'ZOOM\n'
            '  Ctrl/Cmd + / - zoom (centred on the highlight), Z = fit highlight,\n'
            '  A = whole recording, mouse wheel scrolls, Ctrl/Cmd + wheel zooms.\n'
            '  Spectrogram contrast: [ / ] or the slider.\n\n'
            'PLAYBACK & VIDEO\n'
            '  Space plays the highlight; pause keeps your place. L loops (seamless),\n'
            '  1-4 = speed. Video is decoded at full resolution by FFmpeg and drawn at\n'
            '  your screen\'s real pixel density. A highlight you loop is kept in memory,\n'
            '  so later passes are perfectly smooth. Scroll on the video to zoom, drag\n'
            '  to move, double-click / Fit to reset, Rotate if a phone video is sideways.\n'
            '  "Pop out video" gives it its own resizable window.\n\n'
            'EDITING\n'
            '  Double-click a piece to edit its text. Delete removes the selected item.\n'
            '  Right-click (two-finger click) for a menu. Ctrl/Cmd+Z undoes. IPA buttons\n'
            '  insert symbols while typing.\n\n'
            'FILES\n'
            '  Ctrl/Cmd+S saves an Excel file with sheets events, codes and\n'
            '  annotator_data - the same format as the MATLAB version, so files move\n'
            '  freely between the two. Loading checks the table first.')
        d = QtWidgets.QDialog(self)
        d.setWindowTitle(APPNAME + ' - Help')
        v = QtWidgets.QVBoxLayout(d)
        t = QtWidgets.QPlainTextEdit(msg)
        t.setReadOnly(True)
        f = QFont('Menlo' if sys.platform == 'darwin' else 'Consolas')
        f.setStyleHint(QFont.Monospace)
        t.setFont(f)
        v.addWidget(t)
        b = QtWidgets.QPushButton('Close')
        b.clicked.connect(d.accept)
        v.addWidget(b)
        d.resize(640, 620)
        d.exec()

    def show_shortcuts(self):
        rows = [('H', 'ANNOTATING'),
                ('K', 'Highlight new speech + double-click', 'Type the transcript (t-t-table ...)'),
                ('K', 'Double-click highlight on Notes', 'Type notes for it'),
                ('K', 'Highlight a word + Enter', 'Fluent event selected'),
                ('K', 'Then highlight a part + Enter', 'Tick type(s) in the list, Enter'),
                ('K', 'In the list: R P B T / 1-6', 'Tick disfluency / error types'),
                ('K', 'W', 'Whole episode (first part to end of word)'),
                ('K', 'Esc', 'Done with this word'),
                ('K', 'R / P / B / T / E', 'Add a type straight away'),
                ('K', 'N', 'Notes for the selected item'),
                ('K', 'Double-click an item', 'Edit text'),
                ('K', 'Right-click / 2-finger click', 'Menu'),
                ('H', 'COLOURS')] + \
               [('C', d['key'].upper(), '%s - %s' % (d['abbr'], d['name']), d['color']) for d in DISFL] + \
               [('C', 'E', 'error (all types)', ERR_COLOR),
                ('H', 'PLAYBACK'),
                ('K', 'Space', 'Play / pause the highlight'),
                ('K', 'L', 'Loop on / off'),
                ('K', '1 / 2 / 3 / 4', 'Speed 1x / 0.75x / 0.5x / 0.25x'),
                ('K', 'Scroll on the video', 'Zoom (drag = move, double-click = fit)'),
                ('H', 'EDITING & NAVIGATION'),
                ('K', 'Delete', 'Delete the selected item'),
                ('K', 'Ctrl/Cmd + Z / S', 'Undo / save'),
                ('K', 'Left / Right', 'Step one video frame (Shift: 1 s)'),
                ('K', 'Ctrl/Cmd  +  /  -', 'Zoom in / out'),
                ('K', 'Z', 'Zoom to the highlight'),
                ('K', 'A', 'Show the whole recording'),
                ('K', '[  /  ]', 'Spectrogram contrast less / more'),
                ('K', 'Mouse wheel', 'Scroll (Ctrl/Cmd + wheel zooms)'),
                ('K', 'F1', 'Full help')]
        html = ['<table cellpadding="3">']
        for r in rows:
            if r[0] == 'H':
                html.append('<tr><td colspan="2"><br><b>%s</b></td></tr>' % r[1])
            elif r[0] == 'K':
                html.append('<tr><td><b>%s</b></td><td>%s</td></tr>' % (r[1], r[2]))
            else:
                html.append('<tr><td><span style="background:%s">&nbsp;&nbsp;&nbsp;</span> <b>%s</b></td>'
                            '<td>%s</td></tr>' % (qc(r[3]).name(), r[1], r[2]))
        html.append('</table>')
        d = QtWidgets.QDialog(self)
        d.setWindowTitle('Keyboard & mouse shortcuts')
        v = QtWidgets.QVBoxLayout(d)
        t = QtWidgets.QTextBrowser()
        t.setHtml(''.join(html))
        v.addWidget(t)
        b = QtWidgets.QPushButton('Close')
        b.clicked.connect(d.accept)
        v.addWidget(b)
        d.resize(600, 720)
        d.exec()

    def closeEvent(self, e):
        try:
            if not self.confirm_unsaved('closing'):
                e.ignore()
                return
        except Exception:
            pass
        self.stop_play()
        if self.worker is not None:
            self.worker.stop()
        if self.aeng is not None:
            self.aeng.close()
        if self.vsrc is not None:
            self.vsrc.close()
        if self.pop is not None:
            self.pop.closeEvent = lambda ev: ev.accept()
            self.pop.close()
        e.accept()


# ----------------------------------------------------------------------------
# File helpers
# ----------------------------------------------------------------------------
def _xl(v):
    if v is None or (isinstance(v, float) and math.isnan(v)) or v == '':
        return None
    return v


def write_xlsx(out, sheets):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append([_xl(v) for v in r])
    wb.save(out)


def write_csv(out, rows):
    with open(out, 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
        for r in rows:
            w.writerow(['' if v is None else ('%.10g' % v if isinstance(v, float) else v) for v in r])


def _read_csv(f):
    with open(f, newline='', encoding='utf-8-sig') as fh:
        rows = [r for r in csv.reader(fh)]
    out = []
    for i, r in enumerate(rows):
        if i == 0:
            out.append(r)
            continue
        conv = []
        for v in r:
            try:
                conv.append(float(v))
            except ValueError:
                conv.append(v)
        out.append(conv)
    return out


def read_saved_table(f):
    """Reopen: the annotator_data sheet (new files) or the old layouts."""
    d, b = os.path.split(f)
    stem, ext = os.path.splitext(b)
    ext = ext.lower()
    if ext == '.csv':
        stem = re.sub(r'_(annotator|codes)$', '', stem)
        alt = os.path.join(d, stem + '_annotator.csv')
        return _read_csv(alt if os.path.exists(alt) else f)
    if ext in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb['annotator_data'] if 'annotator_data' in wb.sheetnames else wb.worksheets[0]
        rows = [['' if v is None else v for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise RuntimeError('No reader for %s files - save it as .xlsx or .csv.' % ext)


# ----------------------------------------------------------------------------
def main():
    if hasattr(Qt, 'HighDpiScaleFactorRoundingPolicy'):
        QtWidgets.QApplication.setHighDpiScaleFactorRoundingPolicy(
            Qt.HighDpiScaleFactorRoundingPolicy.PassThrough)
    app = QtWidgets.QApplication(sys.argv)
    app.setApplicationName(APPNAME)
    app.setStyle('Fusion')
    pal = app.palette()
    pal.setColor(QtGui.QPalette.Window, qc(BG))
    pal.setColor(QtGui.QPalette.WindowText, Qt.black)
    pal.setColor(QtGui.QPalette.Base, Qt.white)
    pal.setColor(QtGui.QPalette.Text, Qt.black)
    pal.setColor(QtGui.QPalette.Button, QColor(230, 230, 230))
    pal.setColor(QtGui.QPalette.ButtonText, Qt.black)
    app.setPalette(pal)
    init = sys.argv[1] if len(sys.argv) > 1 else None
    w = MainWindow(init)
    w.show()
    sys.exit(app.exec())


if __name__ == '__main__':
    main()
